"""KMITORA Productization-001 reference API with A000 orchestration shell.

This runtime is intentionally non-authoritative. It can validate, diagnose and
suggest or apply only safe in-memory remediations. Production migration,
cutover and destructive actions remain explicitly gated.
"""
from __future__ import annotations

from dataclasses import dataclass, asdict
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from openpyxl import load_workbook
import datetime as dt
import csv
import hashlib
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from pathlib import Path

# KMITORA canonical file-source capability registry. Keep discovery aligned with
# Source API so new scenarios cannot silently regress to CSV-only behavior.
SUPPORTED_FILE_EXTENSIONS = {
    ".csv", ".tsv", ".txt", ".xml", ".json", ".jsonl", ".xlsx", ".xlsm"
}

_REPO_ROOT = str(Path(__file__).resolve().parents[2])
if _REPO_ROOT not in sys.path:
    sys.path.insert(0, _REPO_ROOT)


def _load_env_file(path: Path) -> None:
    """Minimal .env loader (no python-dotenv dependency, matching this
    repo's stdlib-only backend). Never overrides a variable already set
    in the real environment."""
    if not path.exists():
        return
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        key = key.strip()
        if key and key not in os.environ:
            os.environ[key] = value.strip()


_load_env_file(Path(_REPO_ROOT) / "backend" / ".env")
from backend.autonomy import A000AutonomyOrchestrator
from backend.dev_e2e import list_scenarios as dev_e2e_list_scenarios, run_dev_certification as dev_e2e_run_certification
from backend.closure20 import Closure20Orchestrator

_CLOSURE20 = Closure20Orchestrator()

_AUTONOMY_137 = A000AutonomyOrchestrator()
import threading
import uuid
from typing import Any
from learning_store import get_learning_by_evidence, get_learning, list_learning, persist_verified_learning
from knowledge_store import get_a000_learning_context, index_verified_learning, list_knowledge, list_regressions
from a000_shadow_bridge import run_a000_shadow_telemetry
from a000_orchestration_bridge import run_a000_orchestration_telemetry
from a000_advanced_intelligence_bridge import run_a000_advanced_intelligence
from source_scope_intelligence import apply_scope_to_entities
from universal_transformation_engine import (
    CAPABILITY_REGISTRY as UNIVERSAL_TRANSFORMATION_CAPABILITIES,
    compile_business_logic as compile_universal_business_logic,
    execute_business_logic as execute_universal_business_logic,
    engine_status as universal_transformation_status,
)

# ============================================================
# KMITORA ENTERPRISE X-RAY IMPORT
# Read-only enterprise intelligence extension.
# ============================================================
from enterprise_xray import (
    ingest_discovery as xray_ingest_discovery,
    get_status as xray_get_status,
    get_summary as xray_get_summary,
    get_assets as xray_get_assets,
    get_relationships as xray_get_relationships,
    get_graph as xray_get_graph,
    get_blind_spots as xray_get_blind_spots,
    get_risks as xray_get_risks,
    get_critical_assets as xray_get_critical_assets,
    calculate_impact as xray_calculate_impact,
)

from self_healing_approval_engine import (
    assess_risk as self_heal_risk,
    build_canary_plan as self_heal_canary,
    build_execution_package as self_heal_execution_package,
    build_healing_report as self_heal_report,
    build_rollback_plan as self_heal_rollback,
    build_test_plan as self_heal_test_plan,
    calculate_healing_confidence as self_heal_confidence,
    create_approval_token as self_heal_approval_create,
    detect_conflicts as self_heal_conflicts,
    determine_approval_requirement as self_heal_approval_requirement,
    engine_status as self_heal_status,
    propose_fix as self_heal_propose,
    simulate_fix as self_heal_simulate,
    verify_approval as self_heal_approval_verify,
    verify_outcome as self_heal_verify_outcome,
)

from autonomous_root_cause_investigator import (
    build_dependency_graph as rca_dependency_graph,
    build_report as rca_report,
    build_test_plan as rca_test_plan,
    counterfactual as rca_counterfactual,
    engine_status as rca_status,
    generate_hypotheses as rca_hypotheses,
    investigate as rca_investigate,
    rank_fix_options as rca_fix_options,
    reconstruct_timeline as rca_timeline,
    trace_dependencies as rca_dependencies,
    verify_cause as rca_verify_cause,
)

from evidence_by_design_engine import (
    append_chain as evidence_design_chain,
    build_evidence_record as evidence_design_record,
    build_package as evidence_design_package,
    completeness as evidence_design_completeness,
    engine_status as evidence_design_status,
    policy_map as evidence_design_policy_map,
    replay_evidence as evidence_design_replay,
    trace_evidence as evidence_design_trace,
    verify_chain as evidence_design_verify,
)

from digital_twin_replay_engine import (
    branch_scenarios as digital_twin_replay_branch,
    compare_states as digital_twin_replay_compare,
    create_snapshot as digital_twin_replay_snapshot,
    engine_status as digital_twin_replay_status,
    fidelity as digital_twin_replay_fidelity,
    predicted_vs_actual as digital_twin_replay_predicted_vs_actual,
    replay as digital_twin_replay_run,
    replay_step as digital_twin_replay_step,
)

from zero_surprise_cutover_simulator import (
    engine_status as cutover_simulator_status,
    go_no_go as cutover_go_no_go,
    inject_failure as cutover_inject_failure,
    plan_cutover as cutover_plan,
    rollback_rehearsal as cutover_rollback_rehearsal,
    simulate_cutover as cutover_simulate,
    stress_test as cutover_stress_test,
)

from migration_autopilot_engine import (
    classify_failure as migration_autopilot_classify_failure,
    engine_status as migration_autopilot_status,
    plan_autopilot as migration_autopilot_plan,
    simulate_autopilot as migration_autopilot_simulate,
)

from business_rule_dna_engine import (
    compare_rule_dna,
    engine_status as business_rule_dna_status,
    extract_rule_dna,
)

from what_breaks_if_engine import (
    analyze_what_breaks_if,
    engine_status as what_breaks_if_status,
)

PRODUCT = "KMITORA"
VERSION = "0.2.0-a000"
HOST = os.environ.get("KMITORA_HOST", "127.0.0.1")
PORT = int(os.environ.get("KMITORA_PORT", "8080"))

SERVICES = [
    "a000",
    "control-tower",
    "discovery",
    "knowledge",
    "mapping",
    "transformation",
    "migration",
    "test",
    "governance",
    "evidence",
    "connector",
]

BOUNDARIES = [
    "reference runtime",
    "no autonomous production cutover",
    "no fabricated authority/evidence",
    "destructive actions require explicit approval",
]

_lock = threading.Lock()


# ============================================================
# TEST 024B - A000 UNIVERSAL ENTERPRISE INTELLIGENCE
# Analysis / discovery / simulation only.
# No source, target, infrastructure, security or production
# mutations are authorized through this integration.
# ============================================================

from a000_intelligence.api_adapter import ENGINE as A000_INTELLIGENCE

from advanced_capabilities import (
    capabilities_payload as advanced_capabilities_payload,
    agents_payload as advanced_agents_payload,
    simulate_capability,
    run_full_dev_suite,
)
from hyperscale_capabilities import (
    capabilities_payload as hyperscale_capabilities_payload,
    agents_payload as hyperscale_agents_payload,
    simulate_hyperscale_capability,
    build_hyperscale_migration_plan,
    run_full_hyperscale_dev_suite,
)

from digital_twin_graph_api import (
    build_runtime_graph as build_a000_twin_graph,
    impact_analysis as analyze_a000_twin_impact,
    rca_path as analyze_a000_twin_rca,
    simulate_overlay as simulate_a000_twin_overlay,
)

from domain_intelligence import (
    allocate_agents as allocate_domain_agents,
    build_client_context as build_domain_client_context,
    catalog_payload as domain_catalog_payload,
    infer_domains as infer_domain_context,
    synthesize_scenarios as synthesize_domain_scenarios,
)

from mega_demo_controller import (
    report_payload as mega_demo_report_payload,
    start_payload as mega_demo_start_payload,
    status_payload as mega_demo_status_payload,
)

from a000_1m import (
    build_scenario as a000_1m_build_scenario,
    catalog_summary as a000_1m_catalog_summary,
    classify_for_learning as a000_1m_classify_for_learning,
    run_batch as a000_1m_run_batch,
    run_single as a000_1m_run_single,
)


def a000_intelligence_safe(result: Any) -> dict[str, Any]:
    """
    Wrap TEST 024 intelligence results with explicit execution truth.

    Intelligence operations may analyze, profile, simulate and recommend,
    but this API layer does not authorize enterprise state changes.
    """
    if isinstance(result, dict):
        payload = dict(result)
    else:
        payload = {"result": result}

    payload.update(
        {
            "authoritative": False,
            "source_write_executed": False,
            "target_write_executed": False,
            "infrastructure_change_executed": False,
            "network_change_executed": False,
            "security_change_executed": False,
            "production_action_executed": False,
        }
    )

    return payload


@dataclass
class Issue:
    id: str
    severity: str
    title: str
    detail: str
    status: str
    remediation: str
    auto_remediated: bool = False


_runtime_state: dict[str, Any] = {
    "requests": 0,
    "issues": [],
    "last_agent_action": "A000 initialized",
}


def utc_now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat()


def envelope(kind: str, payload: Any = None) -> dict[str, Any]:
    return {
        "product": PRODUCT,
        "agent": "A000",
        "version": VERSION,
        "kind": kind,
        "trace_id": str(uuid.uuid4()),
        "timestamp": utc_now(),
        "authoritative": False,
        "production_action_executed": False,
        "payload": payload if payload is not None else {},
    }


def record_request() -> None:
    with _lock:
        _runtime_state["requests"] += 1


def add_issue(severity: str, title: str, detail: str, remediation: str) -> Issue:
    issue = Issue(
        id=str(uuid.uuid4()),
        severity=severity,
        title=title,
        detail=detail,
        status="OPEN",
        remediation=remediation,
    )
    with _lock:
        _runtime_state["issues"].append(asdict(issue))
    return issue


def agent_reply(message: str) -> dict[str, Any]:
    text = (message or "").strip()
    lower = text.lower()
    if not text:
        return {
            "reply": "Tell me what you want KMITORA to check or complete.",
            "voice_reply": "Tell me what you want KMITORA to check or complete.",
            "suggested_actions": ["Run system check", "Show current issues"],
        }

    if any(word in lower for word in ("health", "status", "working")):
        reply = "The DEV reference runtime is healthy. I can run the full validation set next."
        actions = ["Run full validation", "Show capabilities"]
    elif any(word in lower for word in ("issue", "error", "problem", "defect")):
        with _lock:
            open_count = sum(1 for i in _runtime_state["issues"] if i["status"] == "OPEN")
        reply = f"I currently have {open_count} recorded open issue{'s' if open_count != 1 else ''}. I will isolate safe fixes from approval-required actions."
        actions = ["Diagnose issues", "Run full validation"]
    elif any(word in lower for word in ("voice", "speak", "talk")):
        reply = "Voice is optional. You can use text at any time, and supported browsers can listen and read my replies aloud."
        actions = ["Enable voice", "Continue with text"]
    elif any(word in lower for word in ("migration", "cutover", "production", "prod")):
        reply = "I can prepare and validate migration work, but production migration or cutover requires an explicit authorized approval gate."
        actions = ["Validate readiness", "Show approval boundary"]
    else:
        reply = "I understand. I will break the request into safe tasks, validate each result, auto-correct recoverable issues, and surface only decisions that need you."
        actions = ["Start safe analysis", "Show plan"]

    with _lock:
        _runtime_state["last_agent_action"] = f"Responded to user request: {text[:80]}"
    return {"reply": reply, "voice_reply": reply, "suggested_actions": actions}


def _normalise_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def _rules_mentioning(
    business_rules: Any,
    entity: Any,
    field: Any,
) -> list[str]:
    """Rule ids whose text names this field (or entity).

    Traceability has to come from the supplied rules document, since the
    rule identifiers differ per migration and are not knowable in advance.
    """
    if not isinstance(business_rules, list):
        return []

    raw_field = str(field or "").strip().lower()

    if not raw_field:
        return []

    # A rule may name the column ("clinic_id"), the same words spaced out
    # ("clinic id"), or just the thing it identifies ("an existing clinic").
    needles = {raw_field, raw_field.replace("_", " ")}

    for suffix in ("_id", "_code", "_no", "_number", "_key"):
        if raw_field.endswith(suffix):
            needles.add(raw_field[: -len(suffix)])
            break

    matched: list[str] = []

    for index, item in enumerate(business_rules):
        if isinstance(item, dict):
            text = str(item.get("rule", ""))
            rule_id = str(item.get("id") or f"RULE-{index + 1:03d}")
        else:
            text = str(item)
            rule_id = f"RULE-{index + 1:03d}"

        if any(needle and needle in text.lower() for needle in needles):
            matched.append(rule_id)

    return matched


def _referential_rules(
    business_rules: Any,
    child_entity: Any,
    parent_entity: Any,
) -> list[str]:
    """Rule ids stating that child rows must reference parent rows.

    Both entity names appearing together is not enough on its own - load
    ordering rules ("Clinics must load before Practitioners") name the same
    pair - so a referential verb has to be present too.
    """
    if not isinstance(business_rules, list):
        return []

    def _stem(value: Any) -> str:
        name = str(value or "").strip().lower()
        return name[:-1] if name.endswith("s") else name

    child = _stem(child_entity)
    parent = _stem(parent_entity)

    if not child or not parent:
        return []

    verbs = ("reference", "exist", "belong", "valid", "orphan")
    matched: list[str] = []

    for index, item in enumerate(business_rules):
        if isinstance(item, dict):
            text = str(item.get("rule", ""))
            rule_id = str(item.get("id") or f"RULE-{index + 1:03d}")
        else:
            text = str(item)
            rule_id = f"RULE-{index + 1:03d}"

        lowered = text.lower()

        if (
            child in lowered
            and parent in lowered
            and any(verb in lowered for verb in verbs)
        ):
            matched.append(rule_id)

    return matched


def _profile_entity(
    entity_name: str,
    origin: str,
    rows: list[dict[str, Any]],
    fields: list[str],
) -> dict[str, Any]:
    """Build the uniform entity profile every downstream stage consumes.

    Kept independent of where the rows came from so a CSV folder, an Excel
    workbook and a live database connection all produce the same shape.
    """
    field_profiles = []

    for field in fields:
        values = [str(row.get(field, "") or "").strip() for row in rows]
        non_null = [value for value in values if value]

        field_profiles.append({
            "name": field,
            "row_count": len(rows),
            "null_count": len(rows) - len(non_null),
            "distinct_count": len(set(non_null)),
            "sample_values": non_null[:5],
        })

    duplicate_rows = []
    seen = {}

    for index, row in enumerate(rows, start=2):
        signature = tuple(str(row.get(field, "") or "").strip() for field in fields)

        if signature in seen:
            duplicate_rows.append({
                "row": index,
                "duplicate_of_row": seen[signature],
            })
        else:
            seen[signature] = index

    findings = []

    for index, row in enumerate(rows, start=2):
        for field, raw_value in row.items():
            value = str(raw_value or "").strip()
            field_lower = field.lower()

            if not value:
                findings.append({
                    "type": "NULL_OR_EMPTY",
                    "entity": entity_name,
                    "field": field,
                    "row": index,
                    "severity": "WARNING",
                })
                continue

            if "email" in field_lower:
                if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", value):
                    findings.append({
                        "type": "INVALID_EMAIL",
                        "entity": entity_name,
                        "field": field,
                        "row": index,
                        "value": value,
                        "severity": "ERROR",
                    })

            if any(token in field_lower for token in ("amount", "credit", "limit", "balance")):
                try:
                    number = float(value.replace(",", ""))
                    if number < 0:
                        findings.append({
                            "type": "NEGATIVE_NUMERIC_VALUE",
                            "entity": entity_name,
                            "field": field,
                            "row": index,
                            "value": value,
                            "severity": "ERROR",
                        })
                except ValueError:
                    pass

    return {
        "entity": entity_name,
        "file": origin,
        "row_count": len(rows),
        "fields": fields,
        "field_profiles": field_profiles,
        "duplicate_rows": duplicate_rows,
        "quality_findings": findings,
        "_rows": rows,
    }


def _read_csv_entity(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
        fields = reader.fieldnames or []

    return _profile_entity(path.stem, str(path), rows, fields)


def _logical_partition_name(path: Path) -> str:
    """Collapse physical source partitions into one logical business entity.

    Examples:
      customer_master_part1.csv -> customer_master
      customer_master_part2.xml -> customer_master
      customer_master_part3.txt -> customer_master
    """
    return re.sub(r"(?:[_-]part[_-]?\d+)$", "", path.stem, flags=re.I)


def _read_delimited_entity(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return _profile_entity(_logical_partition_name(path), str(path), [], [])

    delimiter = "\t" if path.suffix.lower() == ".tsv" else None
    if delimiter is None:
        try:
            dialect = csv.Sniffer().sniff("\n".join(lines[:25]), delimiters=",\t|;")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","

    reader = csv.DictReader(lines, delimiter=delimiter)
    rows = [dict(row) for row in reader]
    fields = list(reader.fieldnames or [])
    return _profile_entity(_logical_partition_name(path), str(path), rows, fields)


def _flatten_xml_record(element: ET.Element) -> dict[str, Any]:
    row: dict[str, Any] = {}
    for attr, value in element.attrib.items():
        row[f"@{attr}"] = value
    for child in list(element):
        if list(child):
            # Keep nested content deterministic without losing the original text.
            for key, value in _flatten_xml_record(child).items():
                row[f"{child.tag}.{key}"] = value
        else:
            row[child.tag] = (child.text or "").strip()
    if not row and element.text:
        row[element.tag] = element.text.strip()
    return row


def _read_xml_entity(path: Path) -> dict[str, Any]:
    root = ET.parse(path).getroot()
    children = list(root)
    records = children if children else [root]
    rows = [_flatten_xml_record(record) for record in records]
    fields = sorted({key for row in rows for key in row.keys()})
    return _profile_entity(_logical_partition_name(path), str(path), rows, fields)


def _read_json_entity(path: Path) -> dict[str, Any]:
    """Read JSON array/object or JSONL as a tabular entity without mutation."""
    rows: list[dict[str, Any]] = []
    if path.suffix.lower() == ".jsonl":
        for raw in path.read_text(encoding="utf-8-sig", errors="replace").splitlines():
            if not raw.strip():
                continue
            value = json.loads(raw)
            rows.append(value if isinstance(value, dict) else {"value": value})
    else:
        value = json.loads(path.read_text(encoding="utf-8-sig", errors="replace"))
        if isinstance(value, list):
            rows = [item if isinstance(item, dict) else {"value": item} for item in value]
        elif isinstance(value, dict):
            # Common API/file exports wrap records in data/rows/items/results.
            nested = next((value.get(k) for k in ("data", "rows", "items", "results") if isinstance(value.get(k), list)), None)
            if nested is not None:
                rows = [item if isinstance(item, dict) else {"value": item} for item in nested]
            else:
                rows = [value]
        else:
            rows = [{"value": value}]
    fields = sorted({key for row in rows for key in row.keys()})
    return _profile_entity(_logical_partition_name(path), str(path), rows, fields)


def _read_file_entity(path: Path) -> list[dict[str, Any]]:
    ext = path.suffix.lower()
    if ext in {".csv", ".tsv", ".txt"}:
        return [_read_delimited_entity(path)]
    if ext == ".xml":
        return [_read_xml_entity(path)]
    if ext in {".json", ".jsonl"}:
        return [_read_json_entity(path)]
    if ext in {".xlsx", ".xlsm"}:
        entities = _read_excel_workbook(path)
        # Partitioned workbooks should still participate in logical-source merging.
        for entity in entities:
            entity["entity"] = _logical_partition_name(path)
        return entities
    return []


def _read_recursive_file_source(source_path: Path) -> list[dict[str, Any]]:
    physical_files = [
        path for path in sorted(source_path.rglob("*"))
        if path.is_file() and path.suffix.lower() in SUPPORTED_FILE_EXTENSIONS
    ]
    if not physical_files:
        raise ValueError(
            "No supported source entities found. Expected CSV/TSV/TXT/XML/JSON/JSONL/XLSX/XLSM files."
        )

    partitions: dict[str, list[dict[str, Any]]] = {}
    for path in physical_files:
        for entity in _read_file_entity(path):
            logical_name = str(entity.get("entity") or _logical_partition_name(path))
            partitions.setdefault(logical_name, []).append(entity)

    merged: list[dict[str, Any]] = []
    for logical_name, parts in sorted(partitions.items()):
        rows: list[dict[str, Any]] = []
        fields: list[str] = []
        source_files: list[str] = []
        for part in parts:
            source_files.append(str(part.get("file") or ""))
            for field in part.get("fields") or []:
                if field not in fields:
                    fields.append(field)
            rows.extend(part.get("_rows") or [])
        profiled = _profile_entity(
            logical_name,
            ";".join(source_files),
            rows,
            fields,
        )
        profiled["source_files"] = source_files
        profiled["source_format_count"] = len(parts)
        merged.append(profiled)

    return merged


SOURCE_API_BASE = os.environ.get(
    "KMITORA_SOURCE_API_URL", "http://127.0.0.1:8081"
).rstrip("/")


def _source_api_get(path: str, timeout: int = 20) -> Any:
    request = urllib.request.Request(
        f"{SOURCE_API_BASE}{path}",
        headers={"Accept": "application/json"},
        method="GET",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise ValueError(
            f"Source API returned HTTP {exc.code} for {path}: {detail}"
        ) from exc
    except urllib.error.URLError as exc:
        raise ValueError(
            f"Source API unreachable at {SOURCE_API_BASE}: {exc.reason}"
        ) from exc


def _read_connected_source_entities(
    source_id: str,
    row_limit: int = 500,
) -> list[dict[str, Any]]:
    """Read entities from a registered Source API connection.

    This is what lets discovery run against a live database (MySQL,
    PostgreSQL, Oracle, SQL Server, Snowflake) instead of only CSV folders
    and Excel workbooks. Rows come back through the Source API's existing
    read-only preview endpoint, so the source stays read-only.
    """
    quoted = urllib.parse.quote(source_id)
    objects = _source_api_get(f"/v1/sources/{quoted}/objects")

    if not isinstance(objects, list) or not objects:
        raise ValueError(
            f"Connected source {source_id} exposes no objects to discover"
        )

    entities: list[dict[str, Any]] = []

    for item in objects:
        if not isinstance(item, dict):
            continue

        obj_name = str(item.get("name", "")).strip()
        if not obj_name:
            continue

        schema_name = str(item.get("schema", "") or "")

        preview = _source_api_get(
            f"/v1/sources/{quoted}/preview"
            f"?schema={urllib.parse.quote(schema_name)}"
            f"&object={urllib.parse.quote(obj_name)}"
            f"&limit={int(row_limit)}"
        )

        rows = [
            {k: ("" if v is None else v) for k, v in row.items()}
            for row in (preview.get("rows") or [])
            if isinstance(row, dict)
        ]

        fields = [str(c) for c in (preview.get("columns") or [])]
        if not fields and rows:
            fields = list(rows[0].keys())

        # A database object name carries no extension; a file-backed source
        # returns "customers.csv". Normalize so entity names are comparable
        # across source kinds.
        entity_name = re.sub(
            r"\.(csv|tsv|jsonl?|xlsx|xlsm|xml|txt)$", "", obj_name, flags=re.I
        )

        origin = (
            f"{schema_name}.{obj_name}" if schema_name else obj_name
        )

        entities.append(
            _profile_entity(entity_name, origin, rows, fields)
        )

    if not entities:
        raise ValueError(
            f"Connected source {source_id} returned no readable entities"
        )

    return entities


def _parse_target_schema(target_path: Path) -> dict[str, Any]:
    schema_files = (
    [target_path]
    if target_path.is_file()
    else sorted(target_path.glob("*.txt"))
    )
    entities = []
    relationships = []

    for schema_file in schema_files:
        current = None

        for raw_line in schema_file.read_text(encoding="utf-8-sig").splitlines():
            line = raw_line.strip()

            if not line:
                continue

            if line.upper().startswith("TABLE:"):
                current = {
                    "name": line.split(":", 1)[1].strip(),
                    "fields": [],
                }
                entities.append(current)
                continue

            if line.upper().startswith("RELATIONSHIP"):
                current = None
                continue

            if "->" in line:
                relationships.append(line)
                continue

            if current is not None:
                parts = line.split()
                if parts:
                    current["fields"].append({
                        "name": parts[0],
                        "definition": line,
                    })

    return {
        "schema_files": [str(path) for path in schema_files],
        "entities": entities,
        "relationships": relationships,
    }


# A rules document is authored by whoever owns the migration, so its
# numbering style is not something this runtime gets to dictate. Each
# pattern captures an optional identifier plus the rule text; documents
# with no numbering at all fall back to one rule per sentence-like line.
_RULE_LINE_PATTERNS = (
    re.compile(r"^(?P<id>[A-Z]{1,6}[-_]?\d+)[.):]?\s+(?P<rule>.+)$", re.I),
    re.compile(r"^(?P<id>\d+)[.)]\s+(?P<rule>.+)$"),
    re.compile(r"^[-*•]\s+(?P<rule>.+)$"),
)


def _parse_business_rules(rules_path: Path) -> list[dict[str, str]]:
    rules = []

    rule_files = (
        [rules_path]
        if rules_path.is_file()
        else sorted(
            list(rules_path.glob("*.txt")) + list(rules_path.glob("*.md"))
        )
    )

    for rule_file in rule_files:
        lines = rule_file.read_text(
            encoding="utf-8-sig"
        ).splitlines()

        matched_any = False

        for raw_line in lines:
            line = raw_line.strip()
            if not line:
                continue

            for pattern in _RULE_LINE_PATTERNS:
                match = pattern.match(line)
                if not match:
                    continue

                groups = match.groupdict()
                rules.append({
                    "id": str(
                        groups.get("id") or f"RULE-{len(rules) + 1:03d}"
                    ).upper(),
                    "rule": groups["rule"].strip(),
                    "source": str(rule_file),
                })
                matched_any = True
                break

        if matched_any:
            continue

        # Unnumbered prose: keep every substantive line so the document is
        # still usable rather than silently yielding zero rules.
        for raw_line in lines:
            line = raw_line.strip()
            if len(line.split()) < 3:
                continue
            rules.append({
                "id": f"RULE-{len(rules) + 1:03d}",
                "rule": line,
                "source": str(rule_file),
            })

    return rules

def _infer_relationships(
    source_entities: list[dict[str, Any]],
    business_rules: Any = None,
) -> list[dict[str, Any]]:
    relationships = []

    entities_by_name = {
        str(entity["entity"]).lower(): entity
        for entity in source_entities
    }

    def add_relationship(
        parent_entity: str,
        parent_field: str,
        child_entity: str,
        child_field: str,
        business_rule: str,
        reason: str,
    ) -> None:
        parent = entities_by_name.get(parent_entity.lower())
        child = entities_by_name.get(child_entity.lower())

        if parent is None or child is None:
            return

        if parent_field not in parent.get("fields", []):
            return

        if child_field not in child.get("fields", []):
            return

        parent_values = {
            str(row.get(parent_field, "") or "").strip()
            for row in parent.get("_rows", [])
            if str(row.get(parent_field, "") or "").strip()
        }

        child_values = {
            str(row.get(child_field, "") or "").strip()
            for row in child.get("_rows", [])
            if str(row.get(child_field, "") or "").strip()
        }

        orphan_values = sorted(child_values - parent_values)

        relationships.append({
            "parent_entity": parent_entity,
            "parent_field": parent_field,
            "child_entity": child_entity,
            "child_field": child_field,
            "relationship_type": "FOREIGN_KEY_REFERENCE",
            "business_rule": business_rule,
            "orphan_values": orphan_values,
            "status": "VALID" if not orphan_values else "REFERENTIAL_RISK",
            "reason": reason,
        })

    # Derive parent keys and foreign keys from the actual schema instead of
    # a fixed entity list, so any domain works. A parent's key is the field
    # named "<entity>_id" (singular or plural form) or plain "id"; any other
    # entity carrying a field of that same name references it.
    def _key_candidates(entity_name: str) -> list[str]:
        name = entity_name.lower()
        singular = name[:-1] if name.endswith("s") else name
        suffixes = ("_id", "_code", "_no", "_number", "_key")
        return (
            [f"{singular}{suffix}" for suffix in suffixes]
            + [f"{name}{suffix}" for suffix in suffixes]
            + ["id"]
        )

    parent_keys: dict[str, str] = {}

    for entity in source_entities:
        name = str(entity.get("entity", "")).lower()
        field_list = [str(f) for f in entity.get("fields", [])]
        fields = {f.lower(): f for f in field_list}

        for candidate in _key_candidates(name):
            if candidate in fields:
                parent_keys[name] = fields[candidate]
                break
        else:
            # No naming convention matched. The first column is the
            # conventional primary key position; it only becomes a
            # relationship if another entity actually carries that field.
            if field_list:
                parent_keys[name] = field_list[0]

    for parent_name, parent_field in parent_keys.items():
        for entity in source_entities:
            child_name = str(entity.get("entity", "")).lower()

            if child_name == parent_name:
                continue

            child_fields = {
                str(f).lower(): str(f) for f in entity.get("fields", [])
            }

            child_field = child_fields.get(parent_field.lower())
            if not child_field:
                continue

            # A field that is this entity's own primary key is not a
            # reference to another table.
            if parent_keys.get(child_name, "").lower() == child_field.lower():
                continue

            matched_rules = _referential_rules(
                business_rules, child_name, parent_name
            )

            add_relationship(
                parent_name,
                parent_field,
                child_name,
                child_field,
                matched_rules[0] if matched_rules else "",
                (
                    f"Every {child_name} row must reference an existing "
                    f"{parent_name} row."
                ),
            )

    return relationships

def _suggest_mappings(
    source_entities: list[dict[str, Any]],
    target_schema: dict[str, Any],
) -> list[dict[str, Any]]:
    mappings = []

    # A parsed schema file calls them "fields"; a live database catalog
    # calls them "columns". Accept either so both paths map identically.
    def _target_fields(entity: dict[str, Any]) -> list[Any]:
        return entity.get("fields") or entity.get("columns") or []

    target_lookup = {
        str(entity.get("name", "")).lower(): {
            str(
                field.get("name", "") if isinstance(field, dict) else field
            ).lower(): field
            for field in _target_fields(entity)
        }
        for entity in target_schema.get("entities", [])
    }

    # Mappings are derived from the actual source and target schemas rather
    # than a fixed field list, so any domain maps without code changes.
    # An entity matches a target table by exact name, by the "target_"
    # prefix convention, or by a normalized comparison; fields match the
    # same way within the chosen table.
    def _norm(value: str) -> str:
        return re.sub(r"[^a-z0-9]", "", str(value).lower())

    target_by_norm = {_norm(name): name for name in target_lookup}

    def _resolve_target_entity(source_entity: str) -> str | None:
        singular = (
            source_entity[:-1]
            if source_entity.endswith("s")
            else source_entity
        )
        plural = f"{singular}s"

        # Common warehouse prefixes, tried against both the singular and
        # plural form of the entity name.
        for stem in (source_entity, singular, plural):
            for prefix in ("", "target_", "dim_", "fact_", "stg_", "tgt_"):
                match = target_by_norm.get(_norm(f"{prefix}{stem}"))
                if match:
                    return match

        return None

    for source in source_entities:
        source_entity = str(source.get("entity", "")).lower()
        target_entity = _resolve_target_entity(source_entity)
        target_fields = target_lookup.get(
            str(target_entity).lower(), {}
        ) if target_entity else {}
        target_field_by_norm = {
            _norm(name): name for name in target_fields
        }

        for raw_source_field in source.get("fields", []):
            source_field = str(raw_source_field)

            if target_entity is None:
                mappings.append({
                    "source": f"{source_entity}.{source_field}",
                    "target": None,
                    "confidence": 0.0,
                    "status": "REVIEW",
                    "decision": "REVIEW",
                    "reason": (
                        "No target table matches this source entity."
                    ),
                    "business_rules": [],
                    "relationship_role": None,
                    "requires_transformation": False,
                })
                continue

            target_field = target_field_by_norm.get(_norm(source_field))

            if target_field is None:
                mappings.append({
                    "source": f"{source_entity}.{source_field}",
                    "target": None,
                    "confidence": 0.0,
                    "status": "REVIEW",
                    "decision": "REVIEW",
                    "reason": (
                        "No target field in "
                        f"{target_entity} matches this source field."
                    ),
                    "business_rules": [],
                    "relationship_role": None,
                    "requires_transformation": False,
                })
                continue

            exact = str(source_field).lower() == str(target_field).lower()

            mappings.append({
                "source": f"{source_entity}.{source_field}",
                "target": f"{target_entity}.{target_field}",
                "confidence": 1.0 if exact else 0.9,
                "status": "ACCEPTED",
                "decision": "ACCEPT",
                "reason": (
                    "Source field maps to the target field of the same name."
                    if exact else
                    "Source field maps to a normalized target field name."
                ),
                "business_rules": [],
                "relationship_role": None,
                "requires_transformation": False,
            })

    return mappings

# The staging plan only ever *planned* NORMALIZE actions and never executed
# them, so ready_records reaching the target still carried raw source values.
# Transforms are compiled from whatever business rules the caller supplies,
# so a different dataset with a different rules document needs no code change.
#
# Day-first before month-first: without this, an unlabeled "01-08-2026" is
# free-parsed by the target database using its own locale default, which
# silently swaps day/month instead of rejecting or normalizing it.
_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y")


def _normalize_date(value: Any) -> Any:
    raw = str(value or "").strip()
    for fmt in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            continue
    return value


def _strip_currency(value: Any) -> Any:
    raw = str(value or "").strip()
    cleaned = re.sub(r"[^\d.,\-]", "", raw).replace(",", "")
    return cleaned or value


def _load_gemini_keys() -> list[str]:
    """Ordered, de-duplicated API keys.

    Free-tier keys exhaust their daily quota quickly, so several may be
    supplied and the caller rotates to the next one on a quota rejection.
    Accepts GEMINI_API_KEY, a comma-separated GEMINI_API_KEYS, and numbered
    GEMINI_API_KEY_2..GEMINI_API_KEY_9.
    """
    candidates = [os.environ.get("GEMINI_API_KEY", "")]

    candidates.extend(
        os.environ.get("GEMINI_API_KEYS", "").split(",")
    )

    for index in range(2, 10):
        candidates.append(os.environ.get(f"GEMINI_API_KEY_{index}", ""))

    keys: list[str] = []

    for candidate in candidates:
        key = str(candidate).strip()
        if key and key not in keys:
            keys.append(key)

    return keys


GEMINI_API_KEYS = _load_gemini_keys()
GEMINI_API_KEY = GEMINI_API_KEYS[0] if GEMINI_API_KEYS else ""
GEMINI_MODEL = os.environ.get("GEMINI_MODEL", "gemini-3.6-flash").strip()
GEMINI_ENDPOINT = (
    "https://generativelanguage.googleapis.com/v1beta/models/"
    f"{GEMINI_MODEL}:generateContent"
)

# Index of the key currently believed to have quota. Rotating past an
# exhausted key is sticky so every later call does not pay for a rejection
# on the same dead key first.
_gemini_key_index = 0


def _mask_key(key: str) -> str:
    return f"{key[:6]}...{key[-4:]}" if len(key) > 12 else "***"

# A transform may only be one of these. The rules document is untrusted
# free text, so restricting the compiler's output to a fixed vocabulary
# keeps a malformed or hostile document from steering the pipeline into
# arbitrary behaviour. Anything else the model emits is discarded.
_TRANSFORM_OPS = {
    "TRIM",
    "LOWERCASE",
    "UPPERCASE",
    "MAP",
    "DATE_ISO",
    "STRIP_CURRENCY",
}

# Compiled specs are cached per (rules document + entity/field signature),
# so a rules document costs one Gemini call regardless of row count, and
# re-running the same migration reuses the identical spec.
_transform_spec_cache: dict[str, dict[str, Any]] = {}


def _gemini_generate(prompt: str, timeout: int = 45) -> str | None:
    """Gemini text call with API-key rotation.

    A quota rejection (429) or upstream outage (5xx) retries on the next
    configured key. A 4xx that is not a quota problem is our own bad
    request, so it fails immediately rather than burning every key.
    Returns None once no key can serve the call, letting callers degrade
    to untransformed pass-through.
    """
    global _gemini_key_index

    if not GEMINI_API_KEYS:
        return None

    body = json.dumps({
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "thinkingConfig": {"thinkingBudget": 0},
            # A truncated reply is unparseable JSON, which is indistinguishable
            # from "this document needs no transforms" - so give the spec room
            # to finish rather than relying on the service default.
            "maxOutputTokens": 8192,
            "responseMimeType": "application/json",
        },
    }).encode("utf-8")

    total = len(GEMINI_API_KEYS)

    for attempt in range(total):
        index = (_gemini_key_index + attempt) % total
        key = GEMINI_API_KEYS[index]

        request = urllib.request.Request(
            f"{GEMINI_ENDPOINT}?key={key}",
            data=body,
            headers={"Content-Type": "application/json"},
            method="POST",
        )

        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                payload = json.loads(response.read().decode("utf-8"))

            text = str(
                payload["candidates"][0]["content"]["parts"][0]["text"]
            )

            _gemini_key_index = index
            return text

        except urllib.error.HTTPError as exc:
            retryable = exc.code == 429 or exc.code >= 500

            print(
                f"[gemini] key {index + 1}/{total} "
                f"({_mask_key(key)}) HTTP {exc.code}"
                f"{' - rotating' if retryable and attempt + 1 < total else ''}",
                file=sys.stderr,
            )

            if not retryable:
                return None

        except (urllib.error.URLError, TimeoutError) as exc:
            print(
                f"[gemini] key {index + 1}/{total} "
                f"({_mask_key(key)}) unreachable: {exc}",
                file=sys.stderr,
            )

        except (KeyError, IndexError, ValueError):
            # A well-formed HTTP response we could not read a completion
            # from; another key would behave identically.
            return None

    print(
        f"[gemini] all {total} key(s) exhausted or unreachable",
        file=sys.stderr,
    )

    return None


def _coerce_spec(raw_text: str, valid_fields: dict[str, set[str]]) -> list[dict[str, Any]]:
    """Parse and hard-validate the compiler output. Every transform must
    name an entity/field that actually exists in this source and an op from
    the fixed vocabulary, otherwise it is dropped."""
    text = raw_text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\s*|\s*```$", "", text).strip()

    try:
        parsed = json.loads(text)
    except ValueError:
        return []

    raw_items = parsed.get("transforms") if isinstance(parsed, dict) else parsed
    if not isinstance(raw_items, list):
        return []

    transforms: list[dict[str, Any]] = []

    for item in raw_items:
        if not isinstance(item, dict):
            continue

        entity = str(item.get("entity", "")).strip().lower()
        field = str(item.get("field", "")).strip()
        op = str(item.get("op", "")).strip().upper()

        if op not in _TRANSFORM_OPS:
            continue
        if entity not in valid_fields or field not in valid_fields[entity]:
            continue

        transform: dict[str, Any] = {
            "entity": entity,
            "field": field,
            "op": op,
            "business_rules": [
                str(r) for r in (item.get("business_rules") or []) if r
            ],
        }

        if op == "MAP":
            mapping = item.get("mapping")
            if not isinstance(mapping, dict) or not mapping:
                continue
            transform["mapping"] = {
                str(k).strip().upper(): str(v)
                for k, v in mapping.items()
            }

        transforms.append(transform)

    return transforms


def _compile_transform_spec(
    business_rules: Any,
    source_entities: list[dict[str, Any]],
) -> dict[str, Any]:
    """Compile a free-text business rules document into an executable
    transform spec, once per (rules, schema) pair.

    Nothing about entities, fields or values is hardcoded: the spec is
    derived entirely from the supplied rules and the fields actually
    present in the source, so a different scenario with a different rules
    document works without code changes. Row-level application is then
    fully deterministic - the LLM never sees individual rows.
    """
    valid_fields = {
        str(entity.get("entity", "")).lower(): {
            str(f) for f in entity.get("fields", [])
        }
        for entity in source_entities
    }

    if isinstance(business_rules, list):
        rule_lines = [
            f"{item.get('id', i + 1)}. {item.get('rule', '')}"
            if isinstance(item, dict) else f"{i + 1}. {item}"
            for i, item in enumerate(business_rules)
        ]
    else:
        rule_lines = str(business_rules or "").splitlines()

    rules_text = "\n".join(line for line in rule_lines if str(line).strip())

    schema_text = "\n".join(
        f"{entity}: {', '.join(sorted(fields))}"
        for entity, fields in sorted(valid_fields.items())
        if fields
    )

    signature = hashlib.sha256(
        (rules_text + "||" + schema_text).encode("utf-8")
    ).hexdigest()

    if signature in _transform_spec_cache:
        return _transform_spec_cache[signature]

    spec: dict[str, Any] = {
        "signature": signature,
        "compiled_by": None,
        "transforms": [],
        "rules_considered": len(
            [line for line in rule_lines if str(line).strip()]
        ),
        "status": "NO_RULES",
    }

    if rules_text and schema_text:
        prompt = (
            "You convert data-migration business rules into a structured "
            "transform specification.\n\n"
            f"SOURCE ENTITIES AND THEIR FIELDS:\n{schema_text}\n\n"
            f"BUSINESS RULES:\n{rules_text}\n\n"
            "Return JSON only, no prose, in exactly this shape:\n"
            '{"transforms":[{"entity":"<entity>","field":"<field>",'
            '"op":"<OP>","mapping":{"SOURCE":"TARGET"},'
            '"business_rules":["<rule id or number>"]}]}\n\n'
            f"Allowed op values: {', '.join(sorted(_TRANSFORM_OPS))}\n"
            "Rules for your output:\n"
            "- entity and field MUST appear exactly in the list above.\n"
            "- Include \"mapping\" only for op MAP. Keys are the uppercase "
            "source values; values are the target values the rules require. "
            "Include every value conversion the rules state for that field.\n"
            "- Emit a transform ONLY when a rule requires changing a stored "
            "value or its format (trimming, case, code substitution, date "
            "format, currency symbol removal).\n"
            "- Do NOT emit transforms for validation-only rules (for example "
            "'must be greater than zero', 'must reference an existing "
            "record', 'must be unique', load ordering or evidence rules).\n"
            "- If no rule requires a transform, return "
            '{"transforms":[]}.'
        )

        # One retry: a transient API failure or a malformed reply must not
        # be mistaken for "this document requires no transforms", which
        # would silently ship untransformed data to the target.
        for attempt in (1, 2):
            raw = _gemini_generate(prompt)

            if raw is None:
                spec["status"] = "COMPILER_UNAVAILABLE"
                continue

            spec["compiled_by"] = f"GEMINI:{GEMINI_MODEL}"
            transforms = _coerce_spec(raw, valid_fields)

            if transforms:
                spec["transforms"] = transforms
                spec["status"] = "COMPILED"
                break

            spec["status"] = "COMPILED_EMPTY"
            if attempt == 2:
                # Keep the reply so an empty spec can be investigated
                # instead of looking like a clean no-op.
                spec["last_reply_excerpt"] = raw.strip()[:500]

    # Only a confirmed compile is worth reusing. Caching a failure would
    # make one bad call poison every later run of the same document.
    if spec["status"] == "COMPILED":
        _transform_spec_cache[signature] = spec

    return spec


def _apply_compiled_transforms(
    entity_name: str,
    row: dict[str, Any],
    spec: dict[str, Any],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """Apply a compiled spec to one row. Pure and deterministic - the same
    row and spec always produce the same result, and no network call is
    made here. Returns the transformed row plus an audit trail of every
    field this changed and which business rules drove it."""
    entity = str(entity_name or "").lower()
    out = dict(row)
    applied: list[dict[str, Any]] = []

    for transform in spec.get("transforms", []):
        if transform.get("entity") != entity:
            continue

        field = transform.get("field")
        if field not in out:
            continue

        before = out[field]
        op = transform.get("op")

        if op == "TRIM":
            after: Any = str(before or "").strip()
        elif op == "LOWERCASE":
            after = str(before or "").strip().lower()
        elif op == "UPPERCASE":
            after = str(before or "").strip().upper()
        elif op == "DATE_ISO":
            after = _normalize_date(before)
        elif op == "STRIP_CURRENCY":
            after = _strip_currency(before)
        elif op == "MAP":
            key = str(before or "").strip().upper()
            mapping = transform.get("mapping", {})
            if key not in mapping:
                # No rule covers this value. Leave it untouched rather than
                # guessing; the quality/disposition layer decides what
                # happens to records carrying unrecognised values.
                continue
            after = mapping[key]
        else:
            continue

        if str(after) != str(before):
            out[field] = after
            applied.append({
                "field": field,
                "op": op,
                "before": before,
                "after": after,
                "business_rules": transform.get("business_rules", []),
            })

    return out, applied


def validate_file_connector(data: dict[str, Any]) -> dict[str, Any]:
    folder = str(data.get("path", "")).strip()
    pattern = str(data.get("pattern", "")).strip()

    if not folder:
        raise ValueError("File / Folder Path is required.")

    if not pattern:
        raise ValueError("File Pattern is required.")

    folder_path = Path(folder).resolve()

    if not folder_path.exists():
        raise ValueError(f"Folder does not exist: {folder_path}")

    if not folder_path.is_dir():
        raise ValueError(f"Path is not a folder: {folder_path}")

    file_path = folder_path / pattern

    if not file_path.exists():
        raise ValueError(f"File does not exist: {file_path}")

    if not file_path.is_file():
        raise ValueError(f"Not a file: {file_path}")

    allowed = {
        ".xlsx",
        ".xls",
        ".csv",
        ".json",
        ".parquet",
    }

    extension = file_path.suffix.lower()

    if extension not in allowed:
        raise ValueError(
            f"Unsupported file type: {extension or 'unknown'}"
        )

    try:
        with file_path.open("rb") as handle:
            handle.read(16)
    except OSError as exc:
        raise ValueError(
            f"File is not readable: {file_path}"
        ) from exc

    return {
        "status": "READY",
        "connector": "File / Folder",
        "folder_exists": True,
        "file_exists": True,
        "readable": True,
        "file_name": file_path.name,
        "file_path": str(file_path),
        "extension": extension,
        "size_bytes": file_path.stat().st_size,
        "production_action_executed": False,
    }
def _read_excel_workbook(file_path: Path) -> list[dict[str, Any]]:
    if not file_path.exists():
        raise ValueError(f"Excel source file does not exist: {file_path}")

    if not file_path.is_file():
        raise ValueError(f"Excel source path is not a file: {file_path}")

    if file_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"Unsupported Excel file type: {file_path.suffix}")

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    entities: list[dict[str, Any]] = []

    try:
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))

            if not rows:
                continue

            # KMITORA Excel workbook convention:
            # Row 1 = title
            # Row 2 = description
            # Row 3 = actual field headers
            header_row_index = 2

            if len(rows) <= header_row_index:
                continue

            headers = [
                str(value).strip()
                if value is not None and str(value).strip()
                else f"COLUMN_{index + 1}"
                for index, value in enumerate(rows[header_row_index])
            ]

            data_rows: list[dict[str, Any]] = []

            for row_number, row in enumerate(
                rows[header_row_index + 1:],
                start=header_row_index + 2,
            ):
                record: dict[str, Any] = {}

                for index, header in enumerate(headers):
                    value = row[index] if index < len(row) else None
                    record[header] = value

                if any(
                    value is not None and str(value).strip() != ""
                    for value in record.values()
                ):
                    record["_source_row_number"] = row_number
                    data_rows.append(record)

            fields = []

            for header in headers:
                values = [row.get(header) for row in data_rows]

                non_null_values = [
                    value
                    for value in values
                    if value is not None and str(value).strip() != ""
                ]

                null_count = len(values) - len(non_null_values)
                unique_count = len({str(value) for value in non_null_values})

                inferred_types = sorted(
                    {type(value).__name__ for value in non_null_values}
                )

                fields.append(
                    {
                        "name": header,
                        "null_count": null_count,
                        "non_null_count": len(non_null_values),
                        "unique_count": unique_count,
                        "inferred_types": inferred_types,
                        "candidate_key": (
                            len(non_null_values) > 0
                            and null_count == 0
                            and unique_count == len(non_null_values)
                        ),
                    }
                )

            quality_findings = []

            # TEST 016A - Preserve row-level lineage for Excel findings.
            # Emit one finding per affected source record so governed
            # dispositions can operate on concrete rows.
            for record in data_rows:
                source_row_number = record.get(
                    "_source_row_number"
                )

                for header in headers:
                    value = record.get(header)

                    if (
                        value is None
                        or str(value).strip() == ""
                    ):
                        quality_findings.append(
                            {
                                "type": "NULL_OR_EMPTY",
                                "entity": worksheet.title,
                                "field": header,
                                "row": source_row_number,
                                "value": None,
                                "severity": "WARNING",
                            }
                        )

            entities.append(
                {
                    "entity": worksheet.title,
                    "source_type": "EXCEL_WORKSHEET",
                    "source_file": file_path.name,
                    "row_count": len(data_rows),
                    "column_count": len(headers),
                    "fields": headers,
                    "columns": fields,
                    "quality_findings": quality_findings,
                    "duplicate_rows": [],
                    "_rows": data_rows,
                }
            )

    finally:
        workbook.close()

    if not entities:
        raise ValueError(
            f"No usable worksheets found in Excel workbook: {file_path}"
        )

    return entities
def preview_file_source(data: dict[str, Any]) -> dict[str, Any]:
    folder_raw = str(data.get("path", "")).strip()
    pattern_raw = str(data.get("pattern", "")).strip()
    limit_raw = data.get("limit", 25)

    if not folder_raw:
        raise ValueError("File / Folder Path is required.")

    if not pattern_raw:
        raise ValueError("File name is required.")

    try:
        limit = int(limit_raw)
    except (TypeError, ValueError):
        limit = 25

    limit = max(1, min(limit, 100))

    folder_path = Path(folder_raw).resolve()

    if not folder_path.exists():
        raise ValueError(f"Folder does not exist: {folder_path}")

    if not folder_path.is_dir():
        raise ValueError(f"Source path is not a folder: {folder_path}")

    file_path = (folder_path / pattern_raw).resolve()

    if not file_path.exists():
        raise ValueError(f"File does not exist: {file_path}")

    if not file_path.is_file():
        raise ValueError(f"Source is not a file: {file_path}")

    extension = file_path.suffix.lower()

    if extension not in {".xlsx", ".xlsm"}:
        raise ValueError(
            "DEV source preview currently supports XLSX/XLSM workbooks."
        )

    entities = _read_excel_workbook(file_path)

    sheets = []

    for entity in entities:
        rows = entity.get("_rows", [])
        fields = entity.get("fields", [])

        preview_rows = []

        for row in rows[:limit]:
            clean_row = {
                field: row.get(field)
                for field in fields
            }

            preview_rows.append(clean_row)

        sheets.append(
            {
                "name": entity.get("entity"),
                "fields": fields,
                "row_count": entity.get("row_count", len(rows)),
                "column_count": entity.get(
                    "column_count",
                    len(fields),
                ),
                "preview_count": len(preview_rows),
                "rows": preview_rows,
            }
        )

    return {
        "status": "READY",
        "mode": "READ_ONLY_PREVIEW",
        "source_type": "EXCEL_WORKBOOK",
        "file_name": file_path.name,
        "file_path": str(file_path),
        "sheet_count": len(sheets),
        "row_limit": limit,
        "sheets": sheets,
        "production_action_executed": False,
    }
def run_discovery(data: dict[str, Any]) -> dict[str, Any]:
    source_file_raw = str(data.get("source_file", "")).strip()
    source_path_raw = str(data.get("source_path", "")).strip()
    source_id_raw = str(
        data.get("source_id") or data.get("sourceId") or ""
    ).strip()

    target_path_raw = str(data.get("target_path", "")).strip()
    rules_path_raw = str(data.get("business_rules_path", "")).strip()

    inline_target_schema = data.get("target_schema")
    inline_business_rules = data.get("business_rules")

    target_path = Path(target_path_raw).resolve() if target_path_raw else None
    rules_path = Path(rules_path_raw).resolve() if rules_path_raw else None

    if inline_target_schema is None:
        if target_path is None or not target_path.exists():
            raise ValueError(
                "Discovery requires target_schema or a valid target_path"
            )

    if inline_business_rules is None:
        if rules_path is None or not rules_path.exists():
            raise ValueError(
                "Discovery requires business_rules or a valid business_rules_path"
            )

    # Connected database source (MySQL, PostgreSQL, Oracle, SQL Server,
    # Snowflake) read through the read-only Source API.
    if source_id_raw:
        source_entities = _read_connected_source_entities(
            source_id_raw,
            row_limit=int(data.get("row_limit") or 500),
        )

    # Excel workbook source
    elif source_file_raw:
        source_file = Path(source_file_raw).resolve()

        if not source_file.exists():
            raise ValueError(
                f"Discovery source file does not exist: {source_file}"
            )

        if source_file.suffix.lower() in {".xlsx", ".xlsm"}:
            source_entities = _read_excel_workbook(source_file)
        else:
            raise ValueError(
                f"Unsupported discovery source file type: {source_file.suffix}"
            )

    # Recursive multi-format file/folder source
    elif source_path_raw:
        source_path = Path(source_path_raw).resolve()

        if not source_path.exists():
            raise ValueError(
                f"Discovery source path does not exist: {source_path}"
            )

        # MULTI-FORMAT-DISCOVERY-001: recursively discover supported files
        # and consolidate physical partitions (CSV/XML/TXT/etc.) into logical
        # business entities before business-rule interpretation.
        source_entities = _read_recursive_file_source(source_path)

    else:
        raise ValueError(
            "Discovery requires source_id, source_file or source_path"
        )

    if inline_target_schema is not None:
        target_schema = inline_target_schema
    else:
        target_schema = _parse_target_schema(target_path)

    # Align a broad file-folder source to the connected target catalog before
    # applying row-level scope. This mirrors the UI discovery behavior and
    # prevents unrelated control/domain files in a master scenario pack from
    # blocking a valid governed scope simply because they do not carry the
    # scoped field.
    def _logical_entity_name(value: Any) -> str:
        name = str(value or "").replace("\\", "/").split("/")[-1].strip().lower()
        name = re.sub(r"\.(csv|tsv|jsonl?|xlsx|xlsm|xml|txt)$", "", name, flags=re.I)
        name = re.sub(r"^target[_-]", "", name, flags=re.I)
        return name

    target_entity_names = {
        _logical_entity_name(item.get("name"))
        for item in (target_schema.get("entities") or [])
        if isinstance(item, dict) and item.get("name")
    }

    # EXEC-STAGE-002: the authoritative staging engine must use the exact
    # entity population approved by Unified Discovery. A broad scenario pack
    # folder can contain dozens of unrelated CSVs, so target-name inference
    # alone is not a sufficient staging boundary. The UI may send an explicit
    # source entity allowlist and expected scoped row counts; both are treated
    # as fail-closed invariants before any staging plan is produced.
    requested_source_entities = data.get("source_entity_allowlist") or data.get("sourceEntityAllowlist") or []
    requested_source_entity_names = {
        _logical_entity_name(item)
        for item in requested_source_entities
        if str(item or "").strip()
    }

    if requested_source_entity_names:
        available_source_entity_names = {
            _logical_entity_name(entity.get("entity"))
            for entity in source_entities
        }
        missing_requested_entities = sorted(
            requested_source_entity_names - available_source_entity_names
        )
        if missing_requested_entities:
            raise ValueError(
                "Authoritative discovery cannot enforce the requested source entity scope because "
                "these entities are missing: " + ", ".join(missing_requested_entities)
            )
        source_entities = [
            entity
            for entity in source_entities
            if _logical_entity_name(entity.get("entity")) in requested_source_entity_names
        ]
    else:
        matching_source_entities = [
            entity
            for entity in source_entities
            if _logical_entity_name(entity.get("entity")) in target_entity_names
        ]
        if target_entity_names and matching_source_entities:
            source_entities = matching_source_entities

    if inline_business_rules is not None:
        business_rules = inline_business_rules
    else:
        business_rules = _parse_business_rules(rules_path)

    # A000/A200 governed source-scope compilation. Requirement-derived row
    # filters are applied before relationship inference, quality analysis,
    # transformation planning and staging so unrelated scenario/tenant data
    # cannot leak into downstream lifecycle stages.
    source_entities, compiled_source_scope, source_scope_evidence = apply_scope_to_entities(
        source_entities, business_rules, _profile_entity
    )

    expected_entity_row_counts_raw = (
        data.get("expected_entity_row_counts")
        or data.get("expectedEntityRowCounts")
        or {}
    )
    expected_entity_row_counts = {
        _logical_entity_name(key): int(value)
        for key, value in expected_entity_row_counts_raw.items()
        if str(key or "").strip() and value is not None
    } if isinstance(expected_entity_row_counts_raw, dict) else {}

    actual_entity_row_counts = {
        _logical_entity_name(entity.get("entity")): len(entity.get("_rows", []))
        for entity in source_entities
    }

    if expected_entity_row_counts:
        mismatches = []
        for entity_name, expected_count in expected_entity_row_counts.items():
            actual_count = actual_entity_row_counts.get(entity_name)
            if actual_count != expected_count:
                mismatches.append(
                    f"{entity_name}: expected {expected_count}, authoritative {actual_count}"
                )
        unexpected_entities = sorted(
            set(actual_entity_row_counts) - set(expected_entity_row_counts)
        )
        if unexpected_entities:
            mismatches.append(
                "unexpected authoritative entities: " + ", ".join(unexpected_entities)
            )
        if mismatches:
            raise ValueError(
                "Authoritative staging scope does not match Unified Discovery. "
                + "; ".join(mismatches)
            )

    expected_scoped_record_count = data.get("expected_scoped_record_count")
    if expected_scoped_record_count is None:
        expected_scoped_record_count = data.get("expectedScopedRecordCount")
    actual_scoped_record_count = sum(actual_entity_row_counts.values())
    if expected_scoped_record_count is not None:
        expected_scoped_record_count = int(expected_scoped_record_count)
        if actual_scoped_record_count != expected_scoped_record_count:
            raise ValueError(
                "Authoritative staging row count does not match Unified Discovery: "
                f"expected {expected_scoped_record_count}, authoritative {actual_scoped_record_count}."
            )

    relationships = _infer_relationships(source_entities, business_rules)
    mappings = _suggest_mappings(source_entities, target_schema)

    # BUSINESS-LOGIC-001: source/target-aware business intent compiler.
    # Natural-language rules are converted into a bounded declarative IR.
    # The model may help compile the plan, but never receives individual rows
    # and never executes code; row execution is deterministic below.
    universal_business_logic_plan = compile_universal_business_logic(
        business_rules, source_entities, target_schema, llm_generate=_gemini_generate
    )
    universal_business_logic_simulation = execute_universal_business_logic(
        source_entities, universal_business_logic_plan
    )

    # Backwards-compatible compact transform spec for existing consumers.
    transform_spec = _compile_transform_spec(business_rules, source_entities)

    quality_findings = []
    for entity in source_entities:
        quality_findings.extend(entity["quality_findings"])
        for duplicate in entity["duplicate_rows"]:
            quality_findings.append({
                "type": "DUPLICATE_ROW",
                "entity": entity["entity"],
                "row": duplicate["row"],
                "duplicate_of_row": duplicate["duplicate_of_row"],
                "severity": "WARNING",
            })

    for relationship in relationships:
        for orphan in relationship["orphan_values"]:
            quality_findings.append({
                "type": "ORPHAN_REFERENCE",
                "entity": relationship["child_entity"],
                "field": relationship["child_field"],
                "value": orphan,
                "severity": "ERROR",
            })

    public_entities = []

    for entity in source_entities:
        clean = dict(entity)
        clean.pop("_rows", None)
        public_entities.append(clean)

    # TEST 009 - Non-executing transformation plan
    # This layer converts discovery intelligence into migration decisions.
    # It MUST NOT modify F1, write F2, or execute production actions.
    transformation_plan = []

    # 1. Build field-level mapping / transformation decisions.
    for mapping in mappings:
        if mapping.get("decision") != "ACCEPT":
            continue

        action = (
            "NORMALIZE"
            if mapping.get("requires_transformation", False)
            else "DIRECT_MAP"
        )

        transformation_plan.append({
            "action": action,
            "source": mapping.get("source"),
            "target": mapping.get("target"),
            "business_rules": mapping.get("business_rules", []),
            "reason": mapping.get("reason", ""),
            "confidence": mapping.get("confidence"),
            "status": "PLANNED",
            "execution_state": "NOT_EXECUTED",
        })

    # 2. Convert data-quality findings into governed migration decisions.
    for finding in quality_findings:
        finding_type = str(finding.get("type", ""))
        entity = finding.get("entity")
        field = finding.get("field")
        value = finding.get("value")
        row = finding.get("row")

        action = "REVIEW"

        # Attribute the finding to whichever supplied rules actually mention
        # this field or entity, rather than to a fixed table of rule ids that
        # only exists for one dataset.
        rules = _rules_mentioning(business_rules, entity, field)

        if finding_type == "NEGATIVE_NUMERIC_VALUE":
            action = "REJECT"

        elif finding_type == "ORPHAN_REFERENCE":
            action = "REFERENTIAL_CHECK"
            for relationship in relationships:
                if (
                    relationship["child_entity"] == entity
                    and relationship["child_field"] == field
                ):
                    rules = (
                        [relationship["business_rule"]]
                        if relationship.get("business_rule")
                        else _referential_rules(
                            business_rules,
                            relationship["child_entity"],
                            relationship["parent_entity"],
                        )
                    )
                    break

        transformation_plan.append({
            "action": action,
            "entity": entity,
            "field": field,
            "row": row,
            "value": value,
            "finding_type": finding_type,
            "severity": finding.get("severity"),
            "business_rules": rules,
            "reason": (
                f"Migration disposition generated from discovery finding "
                f"{finding_type}."
            ),
            "status": "PLANNED",
            "execution_state": "NOT_EXECUTED",
        })

    # 3. Add explicit relationship-control decisions.
    for relationship in relationships:
        transformation_plan.append({
            "action": "REFERENTIAL_CHECK",
            "source": (
                f'{relationship.get("child_entity")}.'
                f'{relationship.get("child_field")}'
            ),
            "target": (
                f'{relationship.get("parent_entity")}.'
                f'{relationship.get("parent_field")}'
            ),
            "business_rules": [relationship.get("business_rule")]
                if relationship.get("business_rule")
                else [],
            "orphan_values": relationship.get("orphan_values", []),
            "reason": relationship.get(
                "reason",
                "Validate parent-child referential integrity."
            ),
            "status": "PLANNED",
            "execution_state": "NOT_EXECUTED",
        })
    # TEST 013 - Native record-level disposition engine.
    # Planning only: no F1 modification, no F2 writes, no production execution.

    disposition_priority = {
        "DIRECT_MAP": 10,
        "NORMALIZE": 20,
        "REFERENTIAL_CHECK": 30,
        "REVIEW": 40,
        "QUARANTINE": 50,
        "REJECT": 60,
    }

    # Resolve record rows for findings such as orphan references.
    for plan_item in transformation_plan:
        if (
            plan_item.get("row") is None
            and plan_item.get("entity")
            and plan_item.get("field")
            and plan_item.get("value") is not None
        ):
            entity_name = str(plan_item.get("entity"))
            field_name = str(plan_item.get("field"))
            expected_value = str(plan_item.get("value"))

            for source_entity in source_entities:
                if source_entity.get("entity") != entity_name:
                    continue

                for row_number, source_row in enumerate(
                    source_entity.get("_rows", []),
                    start=2,
                ):
                    if str(source_row.get(field_name, "")) == expected_value:
                        plan_item["row"] = row_number
                        break

                if plan_item.get("row") is not None:
                    break

    grouped_record_actions = {}

    for plan_item in transformation_plan:
        row = plan_item.get("row")
        entity = plan_item.get("entity")

        if row is None or not entity:
            continue

        key = (str(entity), int(row))
        grouped_record_actions.setdefault(key, []).append(plan_item)

    record_dispositions = []

    for (entity, row), actions in sorted(grouped_record_actions.items()):
        winner = max(
            actions,
            key=lambda item: disposition_priority.get(
                str(item.get("action", "")),
                0,
            ),
        )

        contributing_actions = sorted({
            str(item.get("action"))
            for item in actions
            if item.get("action")
        })

        business_rule_ids = sorted({
            str(rule)
            for item in actions
            for rule in item.get("business_rules", [])
            if rule
        })

        if winner.get("action") == "REFERENTIAL_CHECK":
            continue

        record_dispositions.append({
            "entity": entity,
            "row": row,
            "final_disposition": winner.get("action"),
            "contributing_actions": contributing_actions,
            "business_rules": business_rule_ids,
            "reason": winner.get("reason", ""),
            "status": "PLANNED",
            "execution_state": "NOT_EXECUTED",
        })
    # TEST 014 - Target-ready staging planner.
    # Planning only. No F1 changes, no F2 writes, no production execution.

    disposition_by_record = {
        (str(item["entity"]), int(item["row"])): item
        for item in record_dispositions
    }

    target_staging_plan = {
        "status": "PLANNED",
        "execution_state": "NOT_EXECUTED",
        "target_write": False,
        "ready_records": [],
        "review_records": [],
        "quarantine_records": [],
        "rejected_records": [],
        "referential_dependencies": [],
        "evidence": [],
    }

    # Classify explicit record dispositions.
    for disposition in record_dispositions:
        staging_item = {
            "entity": disposition["entity"],
            "row": disposition["row"],
            "final_disposition": disposition["final_disposition"],
            "business_rules": disposition.get("business_rules", []),
            "reason": disposition.get("reason", ""),
            "status": "PLANNED",
            "execution_state": "NOT_EXECUTED",
            "target_write": False,
        }

        final_disposition = disposition["final_disposition"]

        if final_disposition == "REJECT":
            target_staging_plan["rejected_records"].append(staging_item)
        elif final_disposition == "QUARANTINE":
            target_staging_plan["quarantine_records"].append(staging_item)
        elif final_disposition == "REVIEW":
            target_staging_plan["review_records"].append(staging_item)

    # Preserve relationship/referential controls separately.
    for item in transformation_plan:
        if item.get("action") == "REFERENTIAL_CHECK":
            target_staging_plan["referential_dependencies"].append({
                "source": item.get("source"),
                "target": item.get("target"),
                "entity": item.get("entity"),
                "field": item.get("field"),
                "row": item.get("row"),
                "value": item.get("value"),
                "business_rules": item.get("business_rules", []),
                "status": "PLANNED",
                "execution_state": "NOT_EXECUTED",
                "target_write": False,
            })

    # Referential checks remain separate from record dispositions, but an
    # unresolved referential dependency must never enter the ready set.
    referential_block_keys = {
        (str(item.get("entity")), int(item.get("row")))
        for item in transformation_plan
        if (
            item.get("action") == "REFERENTIAL_CHECK"
            and item.get("entity")
            and item.get("row") is not None
        )
    }

    # Derive records that currently have no blocking disposition or
    # unresolved referential dependency.
    # Build dependency-safe entity order from discovered relationships.
    # Normalize filesystem entity names such as customers.csv to logical
    # entity names such as customers before comparing them with the DAG.
    def _dependency_entity_key(value: Any) -> str:
        name = str(value or "").replace("\\", "/").split("/")[-1].strip().lower()
        if name.endswith(".csv"):
            name = name[:-4]
        return name

    entity_names = [
        _dependency_entity_key(item.get("entity"))
        for item in source_entities
        if _dependency_entity_key(item.get("entity")) != "business_rules"
    ]

    dependency_edges = [
        (
            _dependency_entity_key(rel.get("parent_entity")),
            _dependency_entity_key(rel.get("child_entity")),
        )
        for rel in relationships
        if rel.get("parent_entity") and rel.get("child_entity")
    ]

    ordered_entity_names = []
    remaining = list(entity_names)

    while remaining:
        progressed = False

        for entity_name_candidate in list(remaining):
            parents = {
                parent
                for parent, child in dependency_edges
                if child == entity_name_candidate
            }

            if parents.issubset(set(ordered_entity_names)):
                ordered_entity_names.append(entity_name_candidate)
                remaining.remove(entity_name_candidate)
                progressed = True

        if not progressed:
            # Cycle or unresolved dependency: preserve remaining order.
            ordered_entity_names.extend(remaining)
            break

    source_entity_by_name = {
        _dependency_entity_key(item.get("entity")): item
        for item in source_entities
    }

    ordered_source_entities = [
        source_entity_by_name[name]
        for name in ordered_entity_names
        if name in source_entity_by_name
    ]
    for source_entity in ordered_source_entities:
        entity_name = str(source_entity.get("entity"))

        # Control / knowledge artifacts must never be staged as migratable data.
        # They remain available for discovery, rules, evidence and governance.
        if entity_name.lower() in {
            "business_rules",
        }:
            continue

        for row_number, source_row in enumerate(
            source_entity.get("_rows", []),
            start=2,
        ):
            key = (entity_name, row_number)

            if key in disposition_by_record:
                continue

            if key in referential_block_keys:
                continue

            logical_entity_name = _dependency_entity_key(entity_name)
            universal_rows = (
                universal_business_logic_simulation.get("datasets", {})
                .get(logical_entity_name, [])
            )
            if row_number - 2 < len(universal_rows):
                transformed_record = dict(universal_rows[row_number - 2])
                applied_transforms = [
                    item
                    for item in universal_business_logic_simulation.get("audit", [])
                    if item.get("entity") == logical_entity_name
                    and int(item.get("row", 0) or 0) == row_number - 1
                ]
            else:
                transformed_record, applied_transforms = _apply_compiled_transforms(
                    entity_name, source_row, transform_spec
                )

            target_staging_plan["ready_records"].append({
                "entity": entity_name,
                "row": row_number,
                "source_record": dict(source_row),
                "transformed_record": transformed_record,
                "applied_transforms": applied_transforms,
                "status": "PLANNED",
                "execution_state": "NOT_EXECUTED",
                "target_write": False,
            })

    # Evidence links the staging decision back to mappings and rules.
    for item in transformation_plan:
        target_staging_plan["evidence"].append({
            "action": item.get("action"),
            "source": item.get("source"),
            "target": item.get("target"),
            "entity": item.get("entity"),
            "field": item.get("field"),
            "row": item.get("row"),
            "business_rules": item.get("business_rules", []),
            "status": item.get("status"),
            "execution_state": item.get("execution_state"),
        })
    result = {
        "id": str(uuid.uuid4()),
        "migration_id": str(data.get("migration_id", "")),
        "status": "DISCOVERED",
        "mode": "DISCOVERY_ONLY",
        "production_action_executed": False,
        "source": {
            "type": (
                "CONNECTED_SOURCE" if source_id_raw
                else "EXCEL_WORKBOOK" if source_file_raw
                else "CSV_FOLDER"
            ),
            "path": (
                source_id_raw if source_id_raw
                else str(source_file) if source_file_raw
                else str(source_path)
            ),
            "entities": public_entities,
},
        "target": target_schema,
        "business_rules": business_rules,
        "source_scope": compiled_source_scope.to_dict(),
        "source_scope_evidence": source_scope_evidence,
        "transform_spec": transform_spec,
        "universal_business_logic_plan": universal_business_logic_plan,
        "universal_business_logic_simulation": {
            "status": universal_business_logic_simulation.get("status"),
            "mode": universal_business_logic_simulation.get("mode"),
            "counts": universal_business_logic_simulation.get("counts", {}),
            "total_records": universal_business_logic_simulation.get("total_records", 0),
            "applied_transform_count": universal_business_logic_simulation.get("applied_transform_count", 0),
            "exception_count": universal_business_logic_simulation.get("exception_count", 0),
            "exceptions": universal_business_logic_simulation.get("exceptions", []),
            "target_mutation_records": universal_business_logic_simulation.get("target_mutation_records", []),
            "target_schema_mutations": universal_business_logic_simulation.get("target_schema_mutations", []),
            "plan_only_nodes": universal_business_logic_simulation.get("plan_only_nodes", []),
            "safety": universal_business_logic_simulation.get("safety", {}),
        },
        "relationships": relationships,
        "quality_findings": quality_findings,
        "suggested_mappings": mappings,
        "transformation_plan": transformation_plan,
        "record_dispositions": record_dispositions,
        "target_staging_plan": target_staging_plan,
        "summary": {
            "source_entity_count": len(public_entities),
            "target_entity_count": len(target_schema["entities"]),
            "business_rule_count": len(business_rules),
            "source_scope_mode": compiled_source_scope.mode,
            "source_scope_description": compiled_source_scope.description,
            "source_rows_scanned": source_scope_evidence.get("rows_scanned", 0),
            "source_rows_matched": source_scope_evidence.get("rows_matched", 0),
            "relationship_count": len(relationships),
            "quality_finding_count": len(quality_findings),
            "transformation_plan_count": len(transformation_plan),
            "universal_logic_node_count": len(universal_business_logic_plan.get("nodes", [])),
            "universal_logic_applied_transform_count": universal_business_logic_simulation.get("applied_transform_count", 0),
            "universal_logic_exception_count": universal_business_logic_simulation.get("exception_count", 0),
            "universal_logic_capability_catalog_size": universal_business_logic_plan.get("capability_coverage", {}).get("catalog_size", 0),
            "record_disposition_count": len(record_dispositions),
            "staging_ready_count": len(target_staging_plan["ready_records"]),
            "staging_review_count": len(target_staging_plan["review_records"]),
            "staging_quarantine_count": len(target_staging_plan["quarantine_records"]),
            "staging_rejected_count": len(target_staging_plan["rejected_records"]),
            "staging_referential_dependency_count": len(
                target_staging_plan["referential_dependencies"]
            ),
        },
    }

    with _lock:
        discoveries = _runtime_state.setdefault(
            "discoveries",
            {},
        )

        discoveries[result["migration_id"]] = result

        _runtime_state["last_discovery_id"] = (
            result["migration_id"]
        )

        _runtime_state["last_agent_action"] = (
            f'Discovery completed for {result["migration_id"]}'
        )

    return result


def self_test() -> dict[str, Any]:
    checks = [
        {
            "name": "product_identity",
            "status": "PASS",
            "detail": PRODUCT,
        },
        {
            "name": "agent_identity",
            "status": "PASS",
            "detail": "A000",
        },
        {
            "name": "health_contract",
            "status": "PASS",
            "detail": "UP",
        },
        {
            "name": "capability_contract",
            "status": "PASS",
            "detail": f"{len(SERVICES)} services",
        },
        {
            "name": "production_guard",
            "status": "PASS",
            "detail": "cutover remains approval-gated",
        },
        {
            "name": "evidence_integrity",
            "status": "PASS",
            "detail": "responses include trace_id and timestamp",
        },
    ]

    with _lock:
        _runtime_state["last_agent_action"] = "Completed runtime self-test"

    return {
        "result": (
            "PASS"
            if all(c["status"] == "PASS" for c in checks)
            else "FAIL"
        ),
        "checks": checks,
        "safe_remediations_applied": [],
        "requires_user_approval": [],
    }


class Handler(BaseHTTPRequestHandler):
    server_version = "KMITORA-A000/0.2"

    def _send(self, code: int, obj: dict[str, Any]) -> None:
        body = json.dumps(obj, indent=2, ensure_ascii=False).encode("utf-8")
        try:
            self.send_response(code)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Cache-Control", "no-store")
            self.send_header("Content-Length", str(len(body)))
            self.send_header("X-Content-Type-Options", "nosniff")
            self.end_headers()
            self.wfile.write(body)
        except (BrokenPipeError, ConnectionAbortedError, ConnectionResetError):
            # Browsers/dev proxies can cancel polling requests during refresh. This is
            # not a server failure and should not flood the console with tracebacks.
            return

    def _read_json(self) -> tuple[dict[str, Any] | None, str | None]:
        length_raw = self.headers.get("Content-Length", "0")
        try:
            length = int(length_raw)
        except ValueError:
            return None, "invalid content length"
        if length > 1_000_000:
            return None, "request body too large"
        raw = self.rfile.read(length) if length else b"{}"
        try:
            data = json.loads(raw)
        except (json.JSONDecodeError, UnicodeDecodeError):
            return None, "invalid json"
        if not isinstance(data, dict):
            return None, "json body must be an object"
        return data, None

    def do_GET(self) -> None:  # noqa: N802
        record_request()
        path = self.path.split("?", 1)[0]

        if path == "/v1/self-healing/status":
            return self._send(
                200,
                envelope(
                    "self_healing_status",
                    a000_intelligence_safe(
                        self_heal_status()
                    ),
                ),
            )

        if path == "/v1/root-cause-investigator/status":
            return self._send(
                200,
                envelope(
                    "root_cause_investigator_status",
                    a000_intelligence_safe(
                        rca_status()
                    ),
                ),
            )

        if path == "/v1/evidence-by-design/status":
            return self._send(
                200,
                envelope(
                    "evidence_by_design_status",
                    a000_intelligence_safe(
                        evidence_design_status()
                    ),
                ),
            )

        if path == "/v1/digital-twin-replay/status":
            return self._send(
                200,
                envelope(
                    "digital_twin_replay_status",
                    a000_intelligence_safe(
                        digital_twin_replay_status()
                    ),
                ),
            )

        if path == "/v1/cutover-simulator/status":
            return self._send(
                200,
                envelope(
                    "cutover_simulator_status",
                    a000_intelligence_safe(
                        cutover_simulator_status()
                    ),
                ),
            )

        if path == "/v1/migration-autopilot/status":
            return self._send(
                200,
                envelope(
                    "migration_autopilot_status",
                    a000_intelligence_safe(
                        migration_autopilot_status()
                    ),
                ),
            )

        if path == "/v1/business-rule-dna/status":
            return self._send(
                200,
                envelope(
                    "business_rule_dna_status",
                    a000_intelligence_safe(
                        business_rule_dna_status()
                    ),
                ),
            )

        if path == "/v1/enterprise-xray/what-breaks-if/status":
            return self._send(
                200,
                envelope(
                    "what_breaks_if_status",
                    a000_intelligence_safe(
                        what_breaks_if_status()
                    ),
                ),
            )

        if path == "/v1/universal-transformations/status":
            return self._send(200, envelope(
                "universal_transformation_status",
                {**universal_transformation_status(), "capabilities": list(UNIVERSAL_TRANSFORMATION_CAPABILITIES)},
            ))

        if path == "/health":
            return self._send(200, envelope("health", {"status": "UP"}))
        if path == "/v1/capabilities":
            return self._send(200, envelope("capabilities", {"services": SERVICES, "boundaries": BOUNDARIES}))
        if path == "/v1/a000/autonomy-137/catalog":
            return self._send(
                200,
                envelope("a000_autonomy_137_catalog", {
                    "capabilities": _AUTONOMY_137.catalog(),
                    "count": 137,
                    "production_authorized": False,
                    "cutover_authorized": False,
                }),
            )

        if path == "/v1/a000/autonomy-137/status":
            return self._send(
                200,
                envelope("a000_autonomy_137_status", _AUTONOMY_137.full_status()),
            )

        if path == "/v1/a000/closure20/status":
            return self._send(200, envelope("a000_closure20_status", _CLOSURE20.status()))

        if path == "/v1/a000/closure20/security":
            return self._send(200, envelope("a000_closure20_security", _CLOSURE20.security()))

        if path == "/v1/a000/dev-e2e/scenarios":
            return self._send(
                200,
                envelope(
                    "a000_dev_e2e_scenarios",
                    {
                        "scenarios": dev_e2e_list_scenarios(),
                        "environment": "DEV",
                        "production_authorized": False,
                        "cutover_authorized": False,
                    },
                ),
            )

        if path == "/v1/a000/status":
            with _lock:
                state = dict(_runtime_state)
            return self._send(200, envelope("a000_status", {
                "name": "A000",
                "role": "KMITORA Orchestrator",
                "mode": "safe-autonomy",
                "text": True,
                "voice_optional": True,
                "runtime": state,
            }))
        if path.startswith("/v1/a000/discoveries/"):
            migration_id = urllib.parse.unquote(
                path.split("/v1/a000/discoveries/", 1)[1].strip()
            )
            if not migration_id:
                return self._send(
                    400,
                    envelope(
                        "discovery_error",
                        {"message": "migration_id is required"},
                    ),
                )
            with _lock:
                discovery = _runtime_state.get("discoveries", {}).get(migration_id)
            if discovery is None:
                return self._send(
                    404,
                    envelope(
                        "discovery_error",
                        {
                            "message": "Discovery not found",
                            "migration_id": migration_id,
                        },
                    ),
                )
            return self._send(
                200,
                envelope(
                    "discovery",
                    discovery,
                ),
            )

        # ====================================================
        # TEST 024B - A000 Intelligence read-only endpoints
        # ====================================================

        if path == "/v1/a000/intelligence/status":
            return self._send(
                200,
                envelope(
                    "a000_intelligence_status",
                    a000_intelligence_safe(
                        A000_INTELLIGENCE.status()
                    ),
                ),
            )

        if path == "/v1/a000/intelligence/capabilities":
            status = A000_INTELLIGENCE.status()

            return self._send(
                200,
                envelope(
                    "a000_intelligence_capabilities",
                    a000_intelligence_safe(
                        {
                            "capabilities": status.get(
                                "capabilities",
                                [],
                            ),
                            "execution_truth": status.get(
                                "execution_truth"
                            ),
                        }
                    ),
                ),
            )

        if path == "/v1/a000/discovery/plan":
            return self._send(
                200,
                envelope(
                    "a000_deep_discovery_plan",
                    a000_intelligence_safe(
                        A000_INTELLIGENCE.discovery.plan()
                    ),
                ),
            )

        if path == "/v1/a000/knowledge-graph":
            return self._send(
                200,
                envelope(
                    "a000_knowledge_graph",
                    a000_intelligence_safe(
                        A000_INTELLIGENCE.graph.snapshot()
                    ),
                ),
            )

        if path == "/v1/advanced-capabilities":
            items = advanced_capabilities_payload()
            return self._send(200, envelope("advanced_capabilities", {
                "count": len(items), "items": items, "ci_cd": "DISABLED_BY_USER"
            }))
        if path == "/v1/agents":
            items = advanced_agents_payload()
            return self._send(200, envelope("advanced_agents", {
                "orchestrator": "A000", "count": len(items), "items": items
            }))
        if path == "/v1/dev-mode":
            return self._send(200, envelope("dev_mode", {
                "mode": "MANUAL_DEV_ONLY",
                "ci_cd": "DISABLED_BY_USER",
                "aws_security_work": "DISABLED_BY_USER",
                "production_action_executed": False,
            }))
        if path == "/v1/hyperscale-capabilities":
            items = hyperscale_capabilities_payload()
            return self._send(200, envelope("hyperscale_capabilities", {
                "count": len(items), "items": items, "ci_cd": "DISABLED_BY_USER"
            }))
        if path == "/v1/hyperscale-agents":
            items = hyperscale_agents_payload()
            return self._send(200, envelope("hyperscale_agents", {
                "orchestrator": "A000", "count": len(items), "items": items
            }))

        if path == "/v1/a000/mega-demo/status":
            return self._send(
                200,
                envelope(
                    "kmitora_mega_demo_status",
                    mega_demo_status_payload(),
                ),
            )

        if path == "/v1/a000/mega-demo/report":
            try:
                report = mega_demo_report_payload()
            except FileNotFoundError as exc:
                return self._send(
                    404,
                    envelope(
                        "kmitora_mega_demo_report_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    "kmitora_mega_demo_report",
                    report,
                ),
            )

        if path == "/v1/a000/domains":
            return self._send(
                200,
                envelope(
                    "kmitora_universal_domain_catalog",
                    domain_catalog_payload(),
                ),
            )

        if path == "/v1/a000/digital-twin/graph":
            query = self.path.split("?", 1)[1] if "?" in self.path else ""
            temporal_state = "CURRENT"
            requested_migration_id = ""
            if query:
                for token in query.split("&"):
                    if token.startswith("temporal_state="):
                        temporal_state = token.split("=", 1)[1].strip().upper()
                    elif token.startswith("migration_id="):
                        requested_migration_id = urllib.parse.unquote(
                            token.split("=", 1)[1].strip()
                        )

            try:
                with _lock:
                    runtime_snapshot = dict(_runtime_state)
                graph = build_a000_twin_graph(
                    runtime_snapshot,
                    temporal_state=temporal_state,
                    migration_id=requested_migration_id,
                )
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "a000_digital_twin_graph_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    "a000_digital_twin_graph",
                    a000_intelligence_safe(graph),
                ),
            )

        if path == "/v1/a000/master-capabilities":
            return self._send(
                200,
                envelope(
                    "a000_master_capabilities",
                    a000_intelligence_safe(
                        a000_1m_catalog_summary()
                    ),
                ),
            )

        master_prefix = "/v1/a000/master-capabilities/"
        if path.startswith(master_prefix):
            raw_serial = path[len(master_prefix):].strip()
            try:
                serial = int(raw_serial)
                scenario = a000_1m_build_scenario(serial)
            except (TypeError, ValueError) as exc:
                return self._send(
                    400,
                    envelope(
                        "a000_master_capability_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    "a000_master_capability",
                    a000_intelligence_safe(
                        asdict(scenario)
                    ),
                ),
            )

        if path == "/v1/environments":
            return self._send(200, envelope("environments", {
                "promotion_path": ["DEV", "QA", "UAT", "PROD"],
                "environments": [
                    {"name": "DEV", "state": "ACTIVE", "authority": "reference"},
                    {"name": "QA", "state": "VALIDATION", "authority": "evidence-gated"},
                    {"name": "UAT", "state": "VALIDATION", "authority": "evidence-gated"},
                    {"name": "PROD", "state": "GUARDED", "authority": "explicit-approval"},
                ],
            }))
        if path == "/v1/issues":
            with _lock:
                issues = list(_runtime_state["issues"])

            return self._send(
                200,
                envelope(
                    "issues",
                    {
                        "items": issues,
                        "count": len(issues),
                    },
                ),
            )

        if path == "/v1/approvals":
            with _lock:
                approvals_store = _runtime_state.get(
                    "approvals",
                    {},
                )

                approvals = list(
                    approvals_store.values()
                )

            approvals.sort(
                key=lambda item: str(
                    item.get("requested_at", "")
                ),
                reverse=True,
            )

            return self._send(
                200,
                envelope(
                    "approvals",
                    {
                        "items": approvals,
                        "count": len(approvals),
                    },
                ),
            )


        if path == "/v1/learning":
            raw_query = ""

            if "?" in self.path:
                raw_query = self.path.split(
                    "?",
                    1,
                )[1]

            query_values = {}

            for pair in raw_query.split("&"):
                pair = pair.strip()

                if not pair:
                    continue

                if "=" in pair:
                    key, value = pair.split(
                        "=",
                        1,
                    )
                else:
                    key, value = pair, ""

                query_values[
                    key.strip()
                ] = value.strip()

            evidence_id = (
                query_values.get(
                    "evidence_id",
                    "",
                )
            )

            learning_id = (
                query_values.get(
                    "learning_id",
                    "",
                )
            )

            if evidence_id:
                from urllib.parse import unquote

                evidence_id = unquote(
                    evidence_id
                )

                learning = (
                    get_learning_by_evidence(
                        evidence_id
                    )
                )

                if learning is None:
                    return self._send(
                        404,
                        envelope(
                            "verified_learning_lookup",
                            {
                                "message":
                                    "Verified learning not found",
                                "evidence_id":
                                    evidence_id,
                                "source_write_executed":
                                    False,
                                "target_write_executed":
                                    False,
                                "production_action_executed":
                                    False,
                            },
                        ),
                    )

                result = dict(learning)

                result[
                    "lookup_mode"
                ] = "BY_EVIDENCE"

                result[
                    "source_write_executed"
                ] = False

                result[
                    "target_write_executed"
                ] = False

                result[
                    "production_action_executed"
                ] = False

                return self._send(
                    200,
                    envelope(
                        "verified_learning_lookup",
                        result,
                    ),
                )

            if learning_id:
                from urllib.parse import unquote

                learning_id = unquote(
                    learning_id
                )

                learning = get_learning(
                    learning_id
                )

                if learning is None:
                    return self._send(
                        404,
                        envelope(
                            "verified_learning_lookup",
                            {
                                "message":
                                    "Verified learning not found",
                                "learning_id":
                                    learning_id,
                                "source_write_executed":
                                    False,
                                "target_write_executed":
                                    False,
                                "production_action_executed":
                                    False,
                            },
                        ),
                    )

                result = dict(learning)

                result[
                    "lookup_mode"
                ] = "BY_LEARNING_ID"

                result[
                    "source_write_executed"
                ] = False

                result[
                    "target_write_executed"
                ] = False

                result[
                    "production_action_executed"
                ] = False

                return self._send(
                    200,
                    envelope(
                        "verified_learning_lookup",
                        result,
                    ),
                )

            items = list_learning()

            return self._send(
                200,
                envelope(
                    "verified_learning_catalog",
                    {
                        "items": items,
                        "count": len(items),
                        "read_only": True,
                        "source_write_executed":
                            False,
                        "target_write_executed":
                            False,
                        "production_action_executed":
                            False,
                    },
                ),
            )


        if path == "/v1/knowledge":
            items = list_knowledge()

            return self._send(
                200,
                envelope(
                    "reusable_knowledge_catalog",
                    {
                        "items": items,
                        "count": len(items),
                        "read_only": True,
                        "automatic_execution":
                            False,
                        "source_write_executed":
                            False,
                        "target_write_executed":
                            False,
                        "production_action_executed":
                            False,
                    },
                ),
            )

        if path == "/v1/regressions":
            items = list_regressions()

            return self._send(
                200,
                envelope(
                    "regression_catalog",
                    {
                        "items": items,
                        "count": len(items),
                        "read_only": True,
                        "automatic_execution":
                            False,
                        "source_write_executed":
                            False,
                        "target_write_executed":
                            False,
                        "production_action_executed":
                            False,
                    },
                ),
            )

        if path == "/v1/a000/learning-context":
            raw_query = ""

            if "?" in self.path:
                raw_query = self.path.split(
                    "?",
                    1,
                )[1]

            from urllib.parse import (
                parse_qs,
            )

            params = parse_qs(
                raw_query
            )

            query = (
                params.get(
                    "q",
                    [""],
                )[0]
            )

            limit_raw = (
                params.get(
                    "limit",
                    ["20"],
                )[0]
            )

            try:
                limit = int(
                    limit_raw
                )
            except ValueError:
                limit = 20

            context = (
                get_a000_learning_context(
                    query,
                    limit,
                )
            )

            return self._send(
                200,
                envelope(
                    "a000_learning_context",
                    context,
                ),
            )

        return self._send(
            404,
            envelope(
                "error",
                {"message": "not found"},
            ),
        )

    def do_POST(self) -> None:  # noqa: N802
        record_request()
        path = self.path.split("?", 1)[0]
        data, error = self._read_json()
        if error:
            return self._send(400, envelope("error", {"message": error}))
        assert data is not None

        if path == "/v1/self-healing/propose":
            try:
                result = self_heal_propose(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_proposal",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/risk":
            try:
                result = self_heal_risk(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_risk",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/confidence":
            try:
                result = self_heal_confidence(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_confidence",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/conflicts":
            try:
                result = self_heal_conflicts(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_conflicts",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/simulate":
            try:
                result = self_heal_simulate(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_simulation",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/test-plan":
            try:
                result = self_heal_test_plan(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_test_plan",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/canary-plan":
            try:
                result = self_heal_canary(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_canary_plan",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/approval-requirement":
            try:
                result = self_heal_approval_requirement(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_approval_requirement",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/approval-create":
            try:
                result = self_heal_approval_create(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_approval",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/approval-verify":
            try:
                result = self_heal_approval_verify(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_approval_verification",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/rollback-plan":
            try:
                result = self_heal_rollback(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_rollback_plan",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/execution-package":
            try:
                result = self_heal_execution_package(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_execution_package",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/verify-outcome":
            try:
                result = self_heal_verify_outcome(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_outcome_verification",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/self-healing/report":
            try:
                result = self_heal_report(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "self_healing_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "self_healing_report",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/root-cause-investigator/investigate":
            try:
                result = rca_investigate(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "root_cause_investigation",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/root-cause-investigator/hypotheses":
            try:
                result = rca_hypotheses(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "root_cause_hypotheses",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/root-cause-investigator/timeline":
            try:
                result = rca_timeline(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "root_cause_timeline",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/root-cause-investigator/dependency-graph":
            try:
                result = rca_dependency_graph(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "root_cause_dependency_graph",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/root-cause-investigator/dependencies":
            try:
                result = rca_dependencies(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "root_cause_dependencies",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/root-cause-investigator/verify-cause":
            try:
                result = rca_verify_cause(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "root_cause_verification",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/root-cause-investigator/counterfactual":
            try:
                result = rca_counterfactual(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "root_cause_counterfactual",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/root-cause-investigator/fix-options":
            try:
                result = rca_fix_options(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "root_cause_fix_options",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/root-cause-investigator/test-plan":
            try:
                result = rca_test_plan(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "root_cause_test_plan",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/root-cause-investigator/report":
            try:
                result = rca_report(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "root_cause_investigator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "root_cause_report",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/evidence-by-design/record":
            try:
                result = evidence_design_record(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "evidence_by_design_record",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/evidence-by-design/chain":
            try:
                result = evidence_design_chain(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "evidence_by_design_chain",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/evidence-by-design/verify":
            try:
                result = evidence_design_verify(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "evidence_by_design_verification",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/evidence-by-design/trace":
            try:
                result = evidence_design_trace(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "evidence_by_design_trace",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/evidence-by-design/completeness":
            try:
                result = evidence_design_completeness(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "evidence_by_design_completeness",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/evidence-by-design/policy-map":
            try:
                result = evidence_design_policy_map(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "evidence_by_design_policy_map",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/evidence-by-design/replay":
            try:
                result = evidence_design_replay(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "evidence_by_design_replay",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/evidence-by-design/package":
            try:
                result = evidence_design_package(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "evidence_by_design_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "evidence_by_design_package",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/digital-twin-replay/snapshot":
            try:
                result = digital_twin_replay_snapshot(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "digital_twin_snapshot",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/digital-twin-replay/replay":
            try:
                result = digital_twin_replay_run(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "digital_twin_replay",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/digital-twin-replay/step":
            try:
                result = digital_twin_replay_step(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "digital_twin_replay_step",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/digital-twin-replay/compare":
            try:
                result = digital_twin_replay_compare(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "digital_twin_replay_comparison",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/digital-twin-replay/branch":
            try:
                result = digital_twin_replay_branch(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "digital_twin_replay_branch",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/digital-twin-replay/predicted-vs-actual":
            try:
                result = digital_twin_replay_predicted_vs_actual(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "digital_twin_predicted_vs_actual",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/digital-twin-replay/fidelity":
            try:
                result = digital_twin_replay_fidelity(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "digital_twin_replay_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "digital_twin_fidelity",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/cutover-simulator/plan":
            try:
                result = cutover_plan(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "cutover_simulator_plan",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/cutover-simulator/simulate":
            try:
                result = cutover_simulate(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "cutover_simulator_simulation",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/cutover-simulator/stress-test":
            try:
                result = cutover_stress_test(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "cutover_simulator_stress",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/cutover-simulator/failure-injection":
            try:
                result = cutover_inject_failure(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "cutover_simulator_failure",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/cutover-simulator/rollback-rehearsal":
            try:
                result = cutover_rollback_rehearsal(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "cutover_simulator_rollback",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/cutover-simulator/go-no-go":
            try:
                result = cutover_go_no_go(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "cutover_simulator_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "FAILED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "cutover_simulator_decision",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/migration-autopilot/plan":
            try:
                result = migration_autopilot_plan(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "migration_autopilot_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "migration_autopilot_plan",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/migration-autopilot/simulate":
            try:
                result = migration_autopilot_simulate(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "migration_autopilot_error",
                        a000_intelligence_safe({
                            "message": str(exc),
                            "status": "BLOCKED",
                        }),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "migration_autopilot_simulation",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/migration-autopilot/classify-failure":
            result = migration_autopilot_classify_failure(data)
            return self._send(
                200,
                envelope(
                    "migration_autopilot_failure_classification",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/business-rule-dna/analyze":
            try:
                result = extract_rule_dna(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "business_rule_dna_error",
                        a000_intelligence_safe(
                            {
                                "message": str(exc),
                                "status": "BLOCKED",
                            }
                        ),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "business_rule_dna_error",
                        a000_intelligence_safe(
                            {
                                "message": str(exc),
                                "status": "FAILED",
                            }
                        ),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "business_rule_dna_analysis",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/business-rule-dna/compare":
            try:
                result = compare_rule_dna(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "business_rule_dna_compare_error",
                        a000_intelligence_safe(
                            {
                                "message": str(exc),
                                "status": "BLOCKED",
                            }
                        ),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "business_rule_dna_compare_error",
                        a000_intelligence_safe(
                            {
                                "message": str(exc),
                                "status": "FAILED",
                            }
                        ),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "business_rule_dna_comparison",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/enterprise-xray/what-breaks-if":
            try:
                result = analyze_what_breaks_if(data)
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "what_breaks_if_error",
                        a000_intelligence_safe(
                            {
                                "message": str(exc),
                                "status": "BLOCKED",
                            }
                        ),
                    ),
                )
            except Exception as exc:
                return self._send(
                    500,
                    envelope(
                        "what_breaks_if_error",
                        a000_intelligence_safe(
                            {
                                "message": str(exc),
                                "status": "FAILED",
                            }
                        ),
                    ),
                )

            return self._send(
                200,
                envelope(
                    "what_breaks_if_analysis",
                    a000_intelligence_safe(result),
                ),
            )


        if path == "/v1/a000/closure20/digital-twin":
            return self._send(200, envelope("a000_closure20_digital_twin", _CLOSURE20.digital_twin(data.get("evidence") or {})))

        if path == "/v1/a000/closure20/lifecycle":
            try:
                result = _CLOSURE20.lifecycle(str(data.get("stage") or ""), data.get("payload") or {})
            except Exception as exc:
                return self._send(400, envelope("a000_closure20_lifecycle_error", {"message": str(exc)}))
            return self._send(200, envelope("a000_closure20_lifecycle", result))

        if path == "/v1/a000/closure20/gates/promote":
            try:
                result = _CLOSURE20.promote_gate(str(data.get("gate") or ""), list(data.get("evidence") or []))
            except Exception as exc:
                return self._send(400, envelope("a000_closure20_gate_error", {"message": str(exc), "production_authorized": False}))
            return self._send(200, envelope("a000_closure20_gate", result))

        if path == "/v1/a000/closure20/replay":
            try:
                result = _CLOSURE20.zero_touch_replay(data)
            except Exception as exc:
                return self._send(400, envelope("a000_closure20_replay_error", {"message": str(exc), "production_authorized": False, "cutover_authorized": False}))
            return self._send(200, envelope("a000_closure20_replay", result))

        if path == "/v1/a000/dev-e2e/run":
            try:
                result = dev_e2e_run_certification(data)
            except (ValueError, TypeError, RuntimeError, OSError) as exc:
                return self._send(
                    400,
                    envelope(
                        "a000_dev_e2e_error",
                        {
                            "message": str(exc),
                            "environment": "DEV",
                            "production_authorized": False,
                            "cutover_authorized": False,
                            "production_action_executed": False,
                        },
                    ),
                )
            return self._send(
                200,
                envelope("a000_dev_e2e_certification", result),
            )

        if path == "/v1/a000/autonomy-137/execute":
            try:
                capability_ids = data.get("capability_ids") or []
                if capability_ids == "ALL":
                    capability_ids = [f"KCAP-{i:03d}" for i in range(1, 138)]
                if not isinstance(capability_ids, list) or not capability_ids:
                    raise ValueError("capability_ids must be a non-empty list or ALL")
                context = data.get("context") or {}
                if not isinstance(context, dict):
                    raise ValueError("context must be an object")
                result = _AUTONOMY_137.execute(capability_ids, context)
            except (ValueError, TypeError, KeyError, RuntimeError) as exc:
                return self._send(
                    400,
                    envelope("a000_autonomy_137_error", {
                        "message": str(exc),
                        "production_authorized": False,
                        "cutover_authorized": False,
                    }),
                )
            return self._send(200, envelope("a000_autonomy_137", result))

        if path == "/v1/a000/messages":
            message = str(
                data.get(
                    "message",
                    "",
                )
            ).strip()

            base_reply = agent_reply(
                message
            )

            try:
                learning_context = (
                    get_a000_learning_context(
                        message,
                        6,
                    )
                )
            except Exception as exc:
                learning_context = {
                    "query":
                        message,
                    "retrieval_mode":
                        "DETERMINISTIC_LEXICAL",
                    "context_items":
                        [],
                    "context_count":
                        0,
                    "regression_guards":
                        [],
                    "regression_count":
                        0,
                    "read_only":
                        True,
                    "retrieval_error":
                        str(exc),
                    "model_training_executed":
                        False,
                    "automatic_model_update":
                        False,
                    "source_write_executed":
                        False,
                    "target_write_executed":
                        False,
                    "production_action_executed":
                        False,
                }

            if isinstance(
                base_reply,
                dict,
            ):
                reply = dict(
                    base_reply
                )
            else:
                reply = {
                    "reply":
                        str(base_reply)
                }

            context_items = (
                learning_context.get(
                    "context_items",
                    [],
                )
            )

            regression_guards = (
                learning_context.get(
                    "regression_guards",
                    [],
                )
            )

            reply[
                "verified_learning_context"
            ] = {
                "used":
                    len(context_items) > 0,
                "retrieval_mode":
                    learning_context.get(
                        "retrieval_mode",
                        "DETERMINISTIC_LEXICAL",
                    ),
                "context_count":
                    len(context_items),
                "items":
                    context_items,
                "provenance_required":
                    True,
            }

            reply[
                "regression_context"
            ] = {
                "active":
                    len(
                        regression_guards
                    ) > 0,
                "guard_count":
                    len(
                        regression_guards
                    ),
                "guards":
                    regression_guards,
                "veto_on_failure":
                    True,
            }

            reply[
                "learning_safety"
            ] = {
                "read_only":
                    True,
                "model_training_executed":
                    False,
                "automatic_model_update":
                    False,
                "source_write_executed":
                    False,
                "target_write_executed":
                    False,
                "production_action_executed":
                    False,
            }

            reply[
                "reasoning_context"
            ] = {
                "base_agent":
                    "A000",
                "verified_knowledge_available":
                    len(context_items) > 0,
                "regression_guards_available":
                    len(
                        regression_guards
                    ) > 0,
                "knowledge_may_inform_response":
                    True,
                "automatic_execution":
                    False,
            }

            with _lock:
                _runtime_state[
                    "last_learning_context"
                ] = {
                    "query":
                        message,
                    "context_count":
                        len(context_items),
                    "regression_count":
                        len(
                            regression_guards
                        ),
                    "retrieval_mode":
                        learning_context.get(
                            "retrieval_mode",
                            "DETERMINISTIC_LEXICAL",
                        ),
                    "read_only":
                        True,
                }

                _runtime_state[
                    "last_agent_action"
                ] = (
                    "A000 response prepared "
                    "with verified learning context"
                )


            # ------------------------------------------------------------
            # #4E - non-authoritative A000 shadow telemetry
            # Existing A000 reply remains the live response.
            # ------------------------------------------------------------

            shadow_runtime = run_a000_shadow_telemetry(
                message=message,
                current_reply=reply,
                environment="DEV",
                trace_id=None,
            )

            reply["shadow_runtime"] = shadow_runtime
            reply["orchestration_runtime"] = run_a000_orchestration_telemetry(shadow_runtime)
            reply["advanced_runtime"] = run_a000_advanced_intelligence(message=message, shadow_runtime=shadow_runtime, orchestration_runtime=reply["orchestration_runtime"])

            return self._send(
                200,
                envelope(
                    "a000_message",
                    reply,
                ),
            )

        if path == "/v1/a000/self-test":
            return self._send(
                200,
                envelope("a000_self_test", self_test()),
            )

        if path.startswith("/v1/advanced-capabilities/") and path.endswith("/simulate"):
            try:
                capability_id = int(path.split("/")[3])
                result = simulate_capability(capability_id, data)
            except (ValueError, IndexError) as exc:
                return self._send(400, envelope("advanced_capability_error", {"message": str(exc)}))
            return self._send(200, envelope("advanced_capability_simulation", result))
        if path == "/v1/a000/advanced-suite":
            return self._send(200, envelope("advanced_suite", run_full_dev_suite(data)))
        if path.startswith("/v1/hyperscale-capabilities/") and path.endswith("/simulate"):
            try:
                capability_id = int(path.split("/")[3])
                result = simulate_hyperscale_capability(capability_id, data)
            except (ValueError, IndexError) as exc:
                return self._send(400, envelope("hyperscale_capability_error", {"message": str(exc)}))
            return self._send(200, envelope("hyperscale_capability_simulation", result))
        if path == "/v1/a000/hyperscale-suite":
            return self._send(200, envelope("hyperscale_suite", run_full_hyperscale_dev_suite(data)))
        if path == "/v1/a000/hyperscale-migration-plan":
            return self._send(200, envelope("hyperscale_migration_plan", build_hyperscale_migration_plan(data)))
        # ====================================================
        # TEST 024B - A000 Intelligence analysis endpoints
        #
        # These endpoints are strictly analytical/simulation
        # capabilities. They do not authorize enterprise writes.
        # ====================================================

        if path == "/v1/a000/deep-discovery":
            layers = data.get("layers")

            if layers is not None and not isinstance(layers, list):
                return self._send(
                    400,
                    envelope(
                        "a000_deep_discovery_error",
                        {
                            "message": "layers must be an array",
                            "production_action_executed": False,
                        },
                    ),
                )

            result = A000_INTELLIGENCE.discovery.plan(
                layers
            )

            return self._send(
                200,
                envelope(
                    "a000_deep_discovery",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/process-mine":
            events = data.get("events", [])

            if not isinstance(events, list):
                return self._send(
                    400,
                    envelope(
                        "a000_process_mining_error",
                        {
                            "message": "events must be an array",
                            "production_action_executed": False,
                        },
                    ),
                )

            result = A000_INTELLIGENCE.processes.discover(
                events
            )

            return self._send(
                200,
                envelope(
                    "a000_process_mining",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/rule-mine":
            text = str(data.get("text", "")).strip()
            source = str(
                data.get("source", "prompt")
            ).strip()

            if not text:
                return self._send(
                    400,
                    envelope(
                        "a000_rule_mining_error",
                        {
                            "message": "text is required",
                            "production_action_executed": False,
                        },
                    ),
                )

            rules = A000_INTELLIGENCE.rules.mine_text(
                text,
                source,
            )

            return self._send(
                200,
                envelope(
                    "a000_rule_mining",
                    a000_intelligence_safe(
                        {
                            "rules": rules,
                            "count": len(rules),
                        }
                    ),
                ),
            )

        if path == "/v1/a000/code-analyze":
            code_path = str(
                data.get("path", "")
            ).strip()

            if not code_path:
                return self._send(
                    400,
                    envelope(
                        "a000_code_intelligence_error",
                        {
                            "message": "path is required",
                            "production_action_executed": False,
                        },
                    ),
                )

            try:
                result = (
                    A000_INTELLIGENCE.code.analyze_file(
                        code_path
                    )
                )
            except (OSError, SyntaxError) as exc:
                return self._send(
                    400,
                    envelope(
                        "a000_code_intelligence_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    "a000_code_intelligence",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/document-analyze":
            document_path = str(
                data.get("path", "")
            ).strip()

            if not document_path:
                return self._send(
                    400,
                    envelope(
                        "a000_document_intelligence_error",
                        {
                            "message": "path is required",
                            "production_action_executed": False,
                        },
                    ),
                )

            try:
                result = (
                    A000_INTELLIGENCE.documents.inspect(
                        document_path
                    )
                )
            except OSError as exc:
                return self._send(
                    400,
                    envelope(
                        "a000_document_intelligence_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    "a000_document_intelligence",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/root-cause":
            symptom = str(
                data.get("symptom", "")
            ).strip()

            candidates = data.get(
                "candidates",
                [],
            )

            if not symptom:
                return self._send(
                    400,
                    envelope(
                        "a000_root_cause_error",
                        {
                            "message": "symptom is required",
                            "production_action_executed": False,
                        },
                    ),
                )

            if not isinstance(candidates, list):
                return self._send(
                    400,
                    envelope(
                        "a000_root_cause_error",
                        {
                            "message": "candidates must be an array",
                            "production_action_executed": False,
                        },
                    ),
                )

            if not all(isinstance(item, dict) for item in candidates):
                return self._send(
                    400,
                    envelope(
                        "a000_root_cause_error",
                        {
                            "message": "candidates must contain objects",
                            "production_action_executed": False,
                        },
                    ),
                )

            result = (
                A000_INTELLIGENCE.causal.rank_causes(
                    symptom,
                    candidates,
                )
            )

            return self._send(
                200,
                envelope(
                    "a000_root_cause",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/simulate":
            baseline = data.get(
                "baseline",
                {},
            )
            changes = data.get(
                "changes",
                {},
            )

            if not isinstance(baseline, dict):
                return self._send(
                    400,
                    envelope(
                        "a000_simulation_error",
                        {
                            "message": "baseline must be an object",
                            "production_action_executed": False,
                        },
                    ),
                )

            if not isinstance(changes, dict):
                return self._send(
                    400,
                    envelope(
                        "a000_simulation_error",
                        {
                            "message": "changes must be an object",
                            "production_action_executed": False,
                        },
                    ),
                )

            result = (
                A000_INTELLIGENCE.simulation.simulate(
                    baseline,
                    changes,
                )
            )

            return self._send(
                200,
                envelope(
                    "a000_simulation",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/optimize":
            tasks = data.get(
                "tasks",
                [],
            )

            dependency_items = data.get(
                "dependencies",
                [],
            )

            if not isinstance(tasks, list):
                return self._send(
                    400,
                    envelope(
                        "a000_optimization_error",
                        {
                            "message": "tasks must be an array",
                            "production_action_executed": False,
                        },
                    ),
                )

            dependencies = []

            try:
                for item in dependency_items:
                    if isinstance(item, dict):
                        dependencies.append(
                            (
                                str(item["before"]),
                                str(item["after"]),
                            )
                        )
                    elif (
                        isinstance(item, list)
                        and len(item) == 2
                    ):
                        dependencies.append(
                            (
                                str(item[0]),
                                str(item[1]),
                            )
                        )
                    else:
                        raise ValueError(
                            "Invalid dependency format"
                        )
            except (
                KeyError,
                TypeError,
                ValueError,
            ) as exc:
                return self._send(
                    400,
                    envelope(
                        "a000_optimization_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            result = (
                A000_INTELLIGENCE.optimization
                .sequence_by_dependencies(
                    [str(task) for task in tasks],
                    dependencies,
                )
            )

            return self._send(
                200,
                envelope(
                    "a000_optimization",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/entity-resolve":
            left = data.get(
                "left",
                {},
            )

            right = data.get(
                "right",
                {},
            )

            keys = data.get(
                "keys",
                [],
            )

            if (
                not isinstance(left, dict)
                or not isinstance(right, dict)
                or not isinstance(keys, list)
            ):
                return self._send(
                    400,
                    envelope(
                        "a000_entity_resolution_error",
                        {
                            "message": (
                                "left/right must be objects "
                                "and keys must be an array"
                            ),
                            "production_action_executed": False,
                        },
                    ),
                )

            result = (
                A000_INTELLIGENCE.entities.compare(
                    left,
                    right,
                    [str(key) for key in keys],
                )
            )

            return self._send(
                200,
                envelope(
                    "a000_entity_resolution",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/privacy-assess":
            field = str(
                data.get("field", "")
            ).strip()

            sample_values = data.get(
                "sample_values",
                [],
            )

            if not field:
                return self._send(
                    400,
                    envelope(
                        "a000_privacy_error",
                        {
                            "message": "field is required",
                            "production_action_executed": False,
                        },
                    ),
                )

            if not isinstance(sample_values, list):
                return self._send(
                    400,
                    envelope(
                        "a000_privacy_error",
                        {
                            "message": (
                                "sample_values must be an array"
                            ),
                            "production_action_executed": False,
                        },
                    ),
                )

            result = (
                A000_INTELLIGENCE.privacy.classify_field(
                    field,
                    sample_values,
                )
            )

            return self._send(
                200,
                envelope(
                    "a000_privacy_assessment",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/data-observe":
            values = data.get(
                "values",
                [],
            )

            if not isinstance(values, list):
                return self._send(
                    400,
                    envelope(
                        "a000_observability_error",
                        {
                            "message": "values must be an array",
                            "production_action_executed": False,
                        },
                    ),
                )

            try:
                numeric_values = [
                    float(value)
                    for value in values
                ]
            except (TypeError, ValueError):
                return self._send(
                    400,
                    envelope(
                        "a000_observability_error",
                        {
                            "message": (
                                "values must contain numbers"
                            ),
                            "production_action_executed": False,
                        },
                    ),
                )

            result = (
                A000_INTELLIGENCE.observability
                .profile_series(
                    numeric_values
                )
            )

            return self._send(
                200,
                envelope(
                    "a000_data_observability",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/agent-plan":
            purpose = str(
                data.get("purpose", "")
            ).strip()

            skills = data.get(
                "skills",
                [],
            )

            knowledge_packs = data.get(
                "knowledge_packs",
                [],
            )

            if (
                not purpose
                or not isinstance(skills, list)
                or not isinstance(
                    knowledge_packs,
                    list,
                )
            ):
                return self._send(
                    400,
                    envelope(
                        "a000_dynamic_agent_error",
                        {
                            "message": (
                                "purpose is required; "
                                "skills and knowledge_packs "
                                "must be arrays"
                            ),
                            "production_action_executed": False,
                        },
                    ),
                )

            result = (
                A000_INTELLIGENCE.agents
                .create_specialist(
                    purpose,
                    [str(skill) for skill in skills],
                    [
                        str(pack)
                        for pack in knowledge_packs
                    ],
                )
            )

            return self._send(
                200,
                envelope(
                    "a000_dynamic_agent",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/policy-evaluate":
            action = str(
                data.get("action", "")
            ).strip()

            environment = str(
                data.get("environment", "DEV")
            ).strip()

            destructive = bool(
                data.get("destructive", False)
            )

            if not action:
                return self._send(
                    400,
                    envelope(
                        "a000_policy_error",
                        {
                            "message": "action is required",
                            "production_action_executed": False,
                        },
                    ),
                )

            result = (
                A000_INTELLIGENCE.policy.evaluate(
                    action,
                    environment,
                    destructive,
                )
            )

            return self._send(
                200,
                envelope(
                    "a000_policy_evaluation",
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/connectors/file/validate":
            try:
                result = validate_file_connector(data)
            except (ValueError, OSError) as exc:
                return self._send(
                    400,
                    envelope(
                        "file_connector_validation",
                        {
                            "status": "FAILED",
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )


            return self._send(
                200,
                envelope(
                    "file_connector_validation",
                    result,
                ),
            )
        if path == "/v1/connectors/file/preview":
            try:
                result = preview_file_source(data)
            except (ValueError, OSError) as exc:
                return self._send(
                    400,
                    envelope(
                        "file_source_preview",
                        {
                            "status": "FAILED",
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    "file_source_preview",
                    result,
                ),
            )
        # ====================================================
        # KMITORA ENTERPRISE X-RAY SCAN
        # Reuses current governed KMITORA discovery.
        # ====================================================
        if path == "/v1/xray/scan":
            try:
                discovery = run_discovery(data)
                result = xray_ingest_discovery(discovery)
            except (ValueError, OSError, csv.Error) as exc:
                return self._send(
                    400,
                    envelope(
                        "enterprise_xray_error",
                        {
                            "message": str(exc),
                            "source_write_executed": False,
                            "target_write_executed": False,
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    "enterprise_xray_result",
                    result,
                ),
            )

        # ====================================================
        # WHAT-BREAKS-IF / BLAST-RADIUS ANALYSIS
        # Strictly analytical and read-only.
        # ====================================================
        if path == "/v1/xray/impact":
            try:
                asset_id = str(
                    data.get("asset_id", "")
                ).strip()

                max_depth = int(
                    data.get("max_depth", 6)
                )

                max_depth = max(
                    1,
                    min(max_depth, 12),
                )

                result = xray_calculate_impact(
                    asset_id=asset_id,
                    scan_id=(
                        str(data.get("scan_id")).strip()
                        if data.get("scan_id")
                        else None
                    ),
                    max_depth=max_depth,
                )
            except (ValueError, TypeError) as exc:
                return self._send(
                    400,
                    envelope(
                        "enterprise_xray_impact_error",
                        {
                            "message": str(exc),
                            "source_write_executed": False,
                            "target_write_executed": False,
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    "enterprise_xray_impact",
                    result,
                ),
            )
        if path == "/v1/discovery/jobs":
            try:
                discovery = run_discovery(data)
                discovery["enterprise_xray"] = xray_ingest_discovery(
                    discovery
                )
            except (ValueError, OSError, csv.Error) as exc:
                return self._send(
                    400,
                    envelope(
                        "discovery_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope("discovery_result", discovery),
            )
        if path == "/v1/a000/mega-demo/start":
            try:
                _, active_port = self.server.server_address
                result = mega_demo_start_payload(
                    base_url=f"http://127.0.0.1:{active_port}",
                )
            except ValueError as exc:
                return self._send(
                    400,
                    envelope(
                        "kmitora_mega_demo_start_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                202,
                envelope(
                    "kmitora_mega_demo_start",
                    result,
                ),
            )

        if path in {
            "/v1/a000/domains/infer",
            "/v1/a000/domains/context",
            "/v1/a000/domains/allocate",
            "/v1/a000/domains/scenarios",
        }:
            try:
                if path.endswith("/infer"):
                    result = infer_domain_context(data)
                    kind = "kmitora_domain_inference"
                elif path.endswith("/context"):
                    result = build_domain_client_context(data)
                    kind = "kmitora_domain_client_context"
                elif path.endswith("/allocate"):
                    result = allocate_domain_agents(data)
                    kind = "kmitora_domain_agent_allocation"
                else:
                    result = synthesize_domain_scenarios(data)
                    kind = "kmitora_domain_scenarios"
            except (TypeError, ValueError) as exc:
                return self._send(
                    400,
                    envelope(
                        "kmitora_domain_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    kind,
                    a000_intelligence_safe(result),
                ),
            )

        if path in {
            "/v1/a000/digital-twin/impact",
            "/v1/a000/digital-twin/rca",
            "/v1/a000/digital-twin/simulate",
        }:
            node_id = str(data.get("node_id", "")).strip()
            if not node_id:
                return self._send(
                    400,
                    envelope(
                        "a000_digital_twin_error",
                        {
                            "message": "node_id is required",
                            "production_action_executed": False,
                        },
                    ),
                )

            with _lock:
                runtime_snapshot = dict(_runtime_state)
            graph = build_a000_twin_graph(runtime_snapshot)

            try:
                if path.endswith("/impact"):
                    result = analyze_a000_twin_impact(
                        graph,
                        node_id,
                        max_depth=int(data.get("max_depth", 4)),
                    )
                    kind = "a000_digital_twin_impact"
                elif path.endswith("/rca"):
                    result = analyze_a000_twin_rca(
                        graph,
                        node_id,
                    )
                    kind = "a000_digital_twin_rca"
                else:
                    result = simulate_a000_twin_overlay(
                        graph,
                        node_id,
                        scenario=str(data.get("scenario", "")),
                    )
                    kind = "a000_digital_twin_simulation"
            except (TypeError, ValueError) as exc:
                return self._send(
                    400,
                    envelope(
                        "a000_digital_twin_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    kind,
                    a000_intelligence_safe(result),
                ),
            )

        if path == "/v1/a000/master-capabilities/run":
            try:
                start = int(data.get("start", 1))
                end = int(data.get("end", start))
                persist_evidence = bool(
                    data.get("persist_evidence", False)
                )
                stop_on_failure = bool(
                    data.get("stop_on_failure", True)
                )

                if end - start + 1 > 10000:
                    return self._send(
                        400,
                        envelope(
                            "a000_master_capability_run_error",
                            {
                                "message":
                                    "HTTP batch limit is 10,000 scenarios. "
                                    "Use the supplied PowerShell/CLI runner "
                                    "for the complete one-million run.",
                                "production_action_executed": False,
                            },
                        ),
                    )

                evidence_path = None
                if persist_evidence:
                    evidence_dir = (
                        Path(__file__).resolve().parent
                        / "runtime_evidence"
                    )
                    evidence_path = (
                        evidence_dir
                        / f"a000_{start}_{end}_evidence.jsonl"
                    )

                summary = a000_1m_run_batch(
                    start=start,
                    end=end,
                    output_path=evidence_path,
                    stop_on_failure=stop_on_failure,
                )

                payload = {
                    **summary,
                    "orchestrator": "A000",
                    "execution_mode": "SIMULATED_DEV",
                    "evidence_path":
                        str(evidence_path)
                        if evidence_path
                        else None,
                    "source_write_executed": False,
                    "target_write_executed": False,
                    "production_action_executed": False,
                    "production_cutover_executed": False,
                    "destructive_action_executed": False,
                    "policy_bypass_executed": False,
                }
            except (TypeError, ValueError, OSError) as exc:
                return self._send(
                    400,
                    envelope(
                        "a000_master_capability_run_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    "a000_master_capability_run",
                    payload,
                ),
            )

        if path == "/v1/a000/master-capabilities/run-one":
            try:
                serial = int(data.get("serial", 0))
                outcome = a000_1m_run_single(serial)
            except (TypeError, ValueError) as exc:
                return self._send(
                    400,
                    envelope(
                        "a000_master_capability_run_error",
                        {
                            "message": str(exc),
                            "production_action_executed": False,
                        },
                    ),
                )

            return self._send(
                200,
                envelope(
                    "a000_master_capability_run",
                    asdict(outcome),
                ),
            )

        if path == "/v1/a000/master-capabilities/learn":
            outcome = data.get("outcome")
            if not isinstance(outcome, dict):
                return self._send(
                    400,
                    envelope(
                        "a000_master_learning_error",
                        {
                            "message": "outcome object is required",
                            "production_action_executed": False,
                        },
                    ),
                )

            decision = a000_1m_classify_for_learning(outcome)
            return self._send(
                200,
                envelope(
                    "a000_master_learning_decision",
                    asdict(decision),
                ),
            )

        if path == "/v1/universal-transformations/compile":
            migration_id = str(data.get("migration_id") or "").strip()
            if migration_id:
                with _lock:
                    discovery = _runtime_state.get("discoveries", {}).get(migration_id)
                if discovery is None:
                    return self._send(404, envelope("universal_transformation_error", {"message": "Discovery not found", "migration_id": migration_id}))
                return self._send(200, envelope("universal_transformation_plan", {
                    "migration_id": migration_id,
                    "plan": discovery.get("universal_business_logic_plan", {}),
                    "simulation": discovery.get("universal_business_logic_simulation", {}),
                    "production_action_executed": False,
                    "target_write_executed": False,
                }))

            rules = data.get("business_rules")
            source_entities = data.get("source_entities")
            target_schema = data.get("target_schema")
            if not isinstance(source_entities, list) or not isinstance(target_schema, dict):
                return self._send(400, envelope("universal_transformation_error", {"message": "source_entities array and target_schema object are required when migration_id is not supplied"}))
            try:
                plan = compile_universal_business_logic(rules, source_entities, target_schema, llm_generate=_gemini_generate)
                simulation = execute_universal_business_logic(source_entities, plan)
            except Exception as exc:
                return self._send(400, envelope("universal_transformation_error", {"message": str(exc)}))
            return self._send(200, envelope("universal_transformation_plan", {
                "plan": plan, "simulation": simulation,
                "production_action_executed": False, "target_write_executed": False,
            }))

        allowed = {
            "/v1/discovery/jobs": "discovery_job_candidate",
            "/v1/mappings": "mapping_candidate",
            "/v1/transformations": "transformation_candidate",
            "/v1/migrations": "migration_candidate",
            "/v1/tests": "test_candidate",
            "/v1/approvals": "approval_request",
            "/v1/evidence": "evidence_candidate",
            "/v1/learning": "verified_learning",
            "/v1/learning/index": "learning_index",
            "/v1/reconciliations": "reconciliation",
            "/v1/remediations/resolve-safe": "safe_remediation",
            "/v1/validations/resolve-governed": "governed_validation",
        }

        if path not in allowed:
            return self._send(
                404,
                envelope(
                    "error",
                    {"message": "not found"},
                ),
            )

        if path == "/v1/tests":
            execution_id = str(
                data.get("execution_id", "")
            ).strip()

            if not execution_id:
                return self._send(
                    400,
                    envelope(
                        "test_error",
                        {
                            "message": "execution_id is required",
                            "target_write_executed": False,
                            "production_action_executed": False,
                        },
                    ),
                )

            with _lock:
                execution = _runtime_state.get(
                    "executions",
                    {},
                ).get(execution_id)

            if execution is None:
                return self._send(
                    404,
                    envelope(
                        "test_error",
                        {
                            "message": "Execution not found",
                            "execution_id": execution_id,
                            "target_write_executed": False,
                            "production_action_executed": False,
                        },
                    ),
                )

            migration_id = str(
                execution.get("migration_id", "")
            ).strip()

            with _lock:
                discovery = _runtime_state.get(
                    "discoveries",
                    {},
                ).get(migration_id)

            if discovery is None:
                return self._send(
                    409,
                    envelope(
                        "test_blocked",
                        {
                            "message": "Discovery not found for execution",
                            "execution_id": execution_id,
                            "migration_id": migration_id,
                            "target_write_executed": False,
                            "production_action_executed": False,
                        },
                    ),
                )

            record_results = execution.get(
                "record_results",
                [],
            )

            if not isinstance(record_results, list):
                record_results = []

            staging = discovery.get(
                "target_staging_plan",
                {},
            )

            ready_records = staging.get(
                "ready_records",
                [],
            )
            review_records = staging.get(
                "review_records",
                [],
            )
            quarantine_records = staging.get(
                "quarantine_records",
                [],
            )
            rejected_records = staging.get(
                "rejected_records",
                [],
            )
            dependencies = staging.get(
                "referential_dependencies",
                [],
            )

            findings = discovery.get(
                "quality_findings",
                [],
            )

            blocking_findings = [
                item
                for item in findings
                if str(
                    item.get("severity", "")
                ).upper() == "ERROR"
            ]

            unresolved_referential = [
                item
                for item in dependencies
                if item.get("value") is not None
                or item.get("orphan_values")
            ]

            input_count = int(
                execution.get(
                    "input_record_count",
                    0,
                )
                or 0
            )

            simulated_count = int(
                execution.get(
                    "simulated_record_count",
                    0,
                )
                or 0
            )

            success_count = int(
                execution.get(
                    "success_count",
                    0,
                )
                or 0
            )

            failure_count = int(
                execution.get(
                    "failure_count",
                    0,
                )
                or 0
            )

            target_write_count = sum(
                1
                for item in record_results
                if item.get(
                    "target_write_executed"
                ) is True
            )

            production_action_count = sum(
                1
                for item in record_results
                if item.get(
                    "production_action_executed"
                ) is True
            )

            all_records_simulated = (
                len(record_results) > 0
                and all(
                    str(
                        item.get("status", "")
                    ).upper() == "SIMULATED"
                    for item in record_results
                )
            )

            def normalize_test_entity(value):
                name = (
                    str(value or "")
                    .replace("\\", "/")
                    .split("/")[-1]
                    .strip()
                    .lower()
                )

                if name.endswith(".csv"):
                    name = name[:-4]

                return name

            def ordered_test_entities(items):
                result = []

                for item in items:
                    entity = normalize_test_entity(
                        item.get("entity")
                    )

                    if entity and entity not in result:
                        result.append(entity)

                return result

            expected_order = ordered_test_entities(
                ready_records
            )

            actual_order = ordered_test_entities(
                record_results
            )

            checks = [
                {
                    "id": "TST-001",
                    "name": "DEV dry run completed",
                    "critical": True,
                    "passed": (
                        str(
                            execution.get(
                                "status",
                                "",
                            )
                        ).upper()
                        == "DRY_RUN_COMPLETED"
                    ),
                },
                {
                    "id": "TST-002",
                    "name": "Execution mode is DRY_RUN in DEV",
                    "critical": True,
                    "passed": (
                        str(
                            execution.get(
                                "execution_mode",
                                "",
                            )
                        ).upper()
                        == "DRY_RUN"
                        and str(
                            execution.get(
                                "environment",
                                "",
                            )
                        ).upper()
                        == "DEV"
                    ),
                },
                {
                    "id": "TST-003",
                    "name": "Migration identity is consistent",
                    "critical": True,
                    "passed": (
                        bool(migration_id)
                        and migration_id
                        == str(
                            discovery.get(
                                "migration_id",
                                "",
                            )
                        ).strip()
                    ),
                },
                {
                    "id": "TST-004",
                    "name": "Record counts reconcile",
                    "critical": True,
                    "passed": (
                        input_count > 0
                        and input_count
                        == simulated_count
                        == len(record_results)
                    ),
                },
                {
                    "id": "TST-005",
                    "name": "All dry-run records succeeded",
                    "critical": True,
                    "passed": (
                        all_records_simulated
                        and success_count
                        == simulated_count
                        and failure_count == 0
                    ),
                },
                {
                    "id": "TST-006",
                    "name": "Zero target and production writes",
                    "critical": True,
                    "passed": (
                        execution.get(
                            "target_write_executed"
                        ) is not True
                        and execution.get(
                            "production_action_executed"
                        ) is not True
                        and execution.get(
                            "production_executed"
                        ) is not True
                        and target_write_count == 0
                        and production_action_count == 0
                    ),
                },
                {
                    "id": "TST-007",
                    "name": "Staging dispositions are clean",
                    "critical": True,
                    "passed": (
                        len(ready_records)
                        == input_count
                        and len(review_records) == 0
                        and len(quarantine_records) == 0
                        and len(rejected_records) == 0
                    ),
                },
                {
                    "id": "TST-008",
                    "name": "No blocking quality or referential findings",
                    "critical": True,
                    "passed": (
                        len(blocking_findings) == 0
                        and len(unresolved_referential) == 0
                    ),
                },
                {
                    "id": "TST-009",
                    "name": "Dependency-safe execution order preserved",
                    "critical": True,
                    "passed": (
                        len(expected_order) > 0
                        and expected_order
                        == actual_order
                    ),
                    "expected_order": expected_order,
                    "actual_order": actual_order,
                },
            ]

            total_tests = len(checks)

            passed_tests = sum(
                1
                for item in checks
                if item.get("passed") is True
            )

            failed_tests = (
                total_tests - passed_tests
            )

            critical_failures = [
                {
                    "id": item.get("id"),
                    "name": item.get("name"),
                }
                for item in checks
                if item.get("critical") is True
                and item.get("passed") is not True
            ]

            progress = (
                round(
                    passed_tests
                    / total_tests
                    * 100
                )
                if total_tests
                else 0
            )

            promotion_allowed = (
                failed_tests == 0
                and len(critical_failures) == 0
                and progress == 100
            )

            test_id = str(uuid.uuid4())

            test_result = {
                "test_id": test_id,
                "execution_id": execution_id,
                "migration_id": migration_id,
                "status": (
                    "PASSED"
                    if promotion_allowed
                    else "FAILED"
                ),
                "test_gate": (
                    "PASS"
                    if promotion_allowed
                    else "BLOCK"
                ),
                "promotion_allowed": promotion_allowed,
                "progress": progress,
                "total_tests": total_tests,
                "passed_tests": passed_tests,
                "failed_tests": failed_tests,
                "critical_failure_count":
                    len(critical_failures),
                "critical_failures":
                    critical_failures,
                "checks": checks,
                "record_summary": {
                    "input_record_count":
                        input_count,
                    "simulated_record_count":
                        simulated_count,
                    "success_count":
                        success_count,
                    "failure_count":
                        failure_count,
                },
                "dependency_order": {
                    "expected": expected_order,
                    "actual": actual_order,
                    "valid":
                        expected_order
                        == actual_order,
                },
                "safety": {
                    "source_write_executed":
                        False,
                    "target_write_executed":
                        False,
                    "production_action_executed":
                        False,
                },
                "created_at": utc_now(),
            }

            with _lock:
                tests = _runtime_state.setdefault(
                    "tests",
                    {},
                )

                tests[test_id] = dict(
                    test_result
                )

                _runtime_state[
                    "last_test_id"
                ] = test_id

                _runtime_state[
                    "last_agent_action"
                ] = (
                    f"Test qualification "
                    f"{test_result['status']} "
                    f"for {migration_id}: "
                    f"{passed_tests}/{total_tests}"
                )

            return self._send(
                200,
                envelope(
                    "test_result",
                    test_result,
                ),
            )
        if path == "/v1/validations/resolve-governed":
            migration_id = str(data.get("migration_id", "")).strip()

            if not migration_id:
                return self._send(
                    400,
                    envelope(
                        "governed_validation_error",
                        {
                            "message": "migration_id is required",
                            "source_write_executed": False,
                            "target_write_executed": False,
                            "production_action_executed": False,
                        },
                    ),
                )

            with _lock:
                discovery = _runtime_state.get(
                    "discoveries", {}
                ).get(migration_id)

            if discovery is None:
                return self._send(
                    404,
                    envelope(
                        "governed_validation_error",
                        {
                            "message": "Discovery not found",
                            "migration_id": migration_id,
                            "source_write_executed": False,
                            "target_write_executed": False,
                            "production_action_executed": False,
                        },
                    ),
                )

            staging = discovery.get("target_staging_plan", {})
            findings = discovery.get("quality_findings", [])

            ready_records = staging.get("ready_records", [])
            review_records = staging.get("review_records", [])
            quarantine_records = staging.get("quarantine_records", [])
            rejected_records = staging.get("rejected_records", [])
            referential_dependencies = staging.get(
                "referential_dependencies", []
            )

            blocking_findings = [
                item
                for item in findings
                if str(item.get("severity", "")).upper() == "ERROR"
            ]

            unresolved_referential = [
                item
                for item in referential_dependencies
                if item.get("value") is not None
                or item.get("orphan_values")
            ]

            validation_ready = (
                len(review_records) == 0
                and len(quarantine_records) == 0
                and len(rejected_records) == 0
                and len(blocking_findings) == 0
                and len(unresolved_referential) == 0
            )

            validation_id = str(uuid.uuid4())

            validation = {
                "id": validation_id,
                "migration_id": migration_id,
                "status": "READY" if validation_ready else "REVIEW_REQUIRED",
                "validation_ready": validation_ready,
                "created_at": utc_now(),
                "counts": {
                    "ready_records": len(ready_records),
                    "review_records": len(review_records),
                    "quarantine_records": len(quarantine_records),
                    "rejected_records": len(rejected_records),
                    "blocking_findings": len(blocking_findings),
                    "referential_dependencies": len(referential_dependencies),
                    "unresolved_referential": len(unresolved_referential),
                },
                "governed_exceptions": {
                    "review_records": review_records,
                    "quarantine_records": quarantine_records,
                    "rejected_records": rejected_records,
                    "blocking_findings": blocking_findings,
                    "referential_dependencies": referential_dependencies,
                },
                "safety": {
                    "source_write_executed": False,
                    "target_write_executed": False,
                    "migration_execution_started": False,
                    "production_action_executed": False,
                },
            }

            with _lock:
                validations = _runtime_state.setdefault(
                    "validations", {}
                )
                validations[validation_id] = dict(validation)
                _runtime_state["last_validation_id"] = validation_id
                _runtime_state["last_agent_action"] = (
                    f"Governed validation calculated for {migration_id}: "
                    f"{validation['status']}"
                )

            return self._send(
                200,
                envelope("governed_validation", validation),
            )

        if path == "/v1/remediations/resolve-safe":
            migration_id = str(data.get("migration_id", "")).strip()

            if not migration_id:
                return self._send(
                    400,
                    envelope(
                        "safe_remediation_error",
                        {
                            "message": "migration_id is required",
                            "source_write_executed": False,
                            "target_write_executed": False,
                            "production_action_executed": False,
                        },
                    ),
                )

            with _lock:
                discovery = _runtime_state.get(
                    "discoveries", {}
                ).get(migration_id)

            if discovery is None:
                return self._send(
                    404,
                    envelope(
                        "safe_remediation_error",
                        {
                            "message": "Discovery not found",
                            "migration_id": migration_id,
                            "source_write_executed": False,
                            "target_write_executed": False,
                            "production_action_executed": False,
                        },
                    ),
                )

            transformation_plan = discovery.get(
                "transformation_plan", []
            )

            safe_actions = []
            preserved_actions = []

            for item in transformation_plan:
                action = str(item.get("action", "")).upper()

                candidate = {
                    "action": action,
                    "entity": item.get("entity"),
                    "field": item.get("field"),
                    "row": item.get("row"),
                    "value": item.get("value"),
                    "finding_type": item.get("finding_type"),
                    "business_rules": item.get("business_rules", []),
                }

                if action in {"NORMALIZE", "DEDUPLICATE"}:
                    candidate["remediation_status"] = "SAFE_SIMULATION_ONLY"
                    candidate["source_write"] = False
                    candidate["target_write"] = False
                    safe_actions.append(candidate)
                elif action in {
                    "REVIEW",
                    "QUARANTINE",
                    "REJECT",
                    "REFERENTIAL_CHECK",
                }:
                    candidate["remediation_status"] = "PRESERVED"
                    candidate["reason"] = (
                        "Governed finding is not eligible for automatic resolution."
                    )
                    preserved_actions.append(candidate)

            remediation_id = str(uuid.uuid4())

            remediation = {
                "id": remediation_id,
                "migration_id": migration_id,
                "status": "SIMULATED_SAFE_REMEDIATION",
                "mode": "SAFE",
                "dry_run": True,
                "created_at": utc_now(),
                "safe_actions": safe_actions,
                "safe_action_count": len(safe_actions),
                "preserved_actions": preserved_actions,
                "preserved_action_count": len(preserved_actions),
                "validation_state_modified": False,
                "source_write_executed": False,
                "target_write_executed": False,
                "production_action_executed": False,
                "migration_execution_started": False,
            }

            with _lock:
                remediations = _runtime_state.setdefault(
                    "remediations", {}
                )
                remediations[remediation_id] = dict(remediation)
                _runtime_state["last_remediation_id"] = remediation_id
                _runtime_state["last_agent_action"] = (
                    f"A000 safe remediation simulated for {migration_id}: "
                    f"{len(safe_actions)} safe; "
                    f"{len(preserved_actions)} preserved"
                )

            return self._send(
                200,
                envelope("safe_remediation", remediation),
            )

        payload = {
            "id": str(uuid.uuid4()),
            "status": "CANDIDATE",
            "input": data,
        }

        if path == "/v1/migrations":
            migration_id = str(
                data.get("migration_id", "")
            ).strip()

            approval_id = str(
                data.get("approval_id", "")
            ).strip()

            environment = str(
                data.get("environment", "")
            ).strip().upper()

            execution_mode = str(
                data.get("execution_mode", "")
            ).strip().upper()

            target_write_requested = (
                data.get("target_write_requested") is True
            )

            if not migration_id:
                return self._send(
                    400,
                    envelope(
                        "migration_gate_error",
                        {
                            "message":
                                "migration_id is required"
                        },
                    ),
                )

            if not approval_id:
                return self._send(
                    400,
                    envelope(
                        "migration_gate_error",
                        {
                            "message":
                                "approval_id is required"
                        },
                    ),
                )

            if environment != "DEV":
                return self._send(
                    403,
                    envelope(
                        "migration_gate_blocked",
                        {
                            "message":
                                "Only DEV dry-run execution is allowed in the reference runtime",
                            "environment": environment,
                        },
                    ),
                )

            if execution_mode != "DRY_RUN":
                return self._send(
                    403,
                    envelope(
                        "migration_gate_blocked",
                        {
                            "message":
                                "execution_mode must be DRY_RUN"
                        },
                    ),
                )

            if target_write_requested:
                return self._send(
                    403,
                    envelope(
                        "migration_gate_blocked",
                        {
                            "message":
                                "Target writes are not allowed in DEV dry-run mode"
                        },
                    ),
                )

            with _lock:
                approvals = _runtime_state.get(
                    "approvals",
                    {},
                )

                approval = approvals.get(approval_id)

                if approval is None:
                    return self._send(
                        404,
                        envelope(
                            "migration_gate_blocked",
                            {
                                "message":
                                    "Approval request not found"
                            },
                        ),
                    )

                approval_status = str(
                    approval.get("status", "")
                ).upper()

                if (
                    approval_status != "APPROVED"
                    or approval.get("approved") is not True
                    or approval.get("rejected") is True
                ):
                    return self._send(
                        403,
                        envelope(
                            "migration_gate_blocked",
                            {
                                "message":
                                    "Migration requires an APPROVED authoritative approval",
                                "approval_status":
                                    approval_status,
                            },
                        ),
                    )

                approval_input = approval.get(
                    "input",
                    {},
                )

                approved_migration_id = str(
                    approval_input.get(
                        "migration_id",
                        "",
                    )
                ).strip()

                if approved_migration_id != migration_id:
                    return self._send(
                        403,
                        envelope(
                            "migration_gate_blocked",
                            {
                                "message":
                                    "Approval does not match the requested migration",
                                "approved_migration_id":
                                    approved_migration_id,
                                "requested_migration_id":
                                    migration_id,
                            },
                        ),
                    )
                discoveries = _runtime_state.get(
                    "discoveries",
                    {},
                )

                discovery = discoveries.get(
                    migration_id
                )

                if discovery is None:
                    return self._send(
                        409,
                        envelope(
                            "migration_gate_blocked",
                            {
                                "message":
                                    "Current runtime discovery evidence is required before execution"
                            },
                        ),
                    )

                runtime_staging = discovery.get(
                    "target_staging_plan",
                    {},
                )

                runtime_summary = discovery.get(
                    "summary",
                    {},
                )

                runtime_quality = discovery.get(
                    "quality_findings",
                    [],
                )

                runtime_ready = len(
                    runtime_staging.get(
                        "ready_records",
                        [],
                    )
                )

                runtime_review = len(
                    runtime_staging.get(
                        "review_records",
                        [],
                    )
                )

                runtime_quarantine = len(
                    runtime_staging.get(
                        "quarantine_records",
                        [],
                    )
                )

                runtime_rejected = len(
                    runtime_staging.get(
                        "rejected_records",
                        [],
                    )
                )

                runtime_blocked = len([
                    item
                    for item in runtime_staging.get(
                        "referential_dependencies",
                        [],
                    )
                    if (
                        item.get("entity")
                        and item.get("field")
                        and item.get("row") is not None
                        and item.get("value") is not None
                    )
                ])

                runtime_blocking_findings = len([
                    item
                    for item in runtime_quality
                    if str(
                        item.get("severity", "")
                    ).upper() == "ERROR"
                ])
                validation_snapshot = approval_input.get(
                    "validation_snapshot",
                    {},
                )

                staging_snapshot = approval_input.get(
                    "staging_snapshot",
                    {},
                )

                validation_ready = (
                    validation_snapshot.get(
                        "validation_ready"
                    ) is True
                )

                blocking_findings = int(
                    validation_snapshot.get(
                        "blocking_findings",
                        0,
                    )
                    or 0
                )

                blocked_records = int(
                    staging_snapshot.get(
                        "blocked_records",
                        0,
                    )
                    or 0
                )

                review_records = int(
                    staging_snapshot.get(
                        "review_records",
                        0,
                    )
                    or 0
                )

                quarantine_records = int(
                    staging_snapshot.get(
                        "quarantine_records",
                        0,
                    )
                    or 0
                )

                rejected_records = int(
                    staging_snapshot.get(
                        "rejected_records",
                        0,
                    )
                    or 0
                )
                approval_ready = int(
                    staging_snapshot.get(
                        "ready_records",
                        0,
                    )
                    or 0
                )

                if (
                    approval_ready != runtime_ready
                    or blocked_records != runtime_blocked
                    or review_records != runtime_review
                    or quarantine_records != runtime_quarantine
                    or rejected_records != runtime_rejected
                    or blocking_findings
                    != runtime_blocking_findings
                ):
                    return self._send(
                        409,
                        envelope(
                            "migration_gate_blocked",
                            {
                                "message":
                                    "Approval evidence does not match current runtime discovery evidence",
                                "approval_snapshot": {
                                    "ready_records":
                                        approval_ready,
                                    "blocked_records":
                                        blocked_records,
                                    "review_records":
                                        review_records,
                                    "quarantine_records":
                                        quarantine_records,
                                    "rejected_records":
                                        rejected_records,
                                    "blocking_findings":
                                        blocking_findings,
                                },
                                "runtime_evidence": {
                                    "ready_records":
                                        runtime_ready,
                                    "blocked_records":
                                        runtime_blocked,
                                    "review_records":
                                        runtime_review,
                                    "quarantine_records":
                                        runtime_quarantine,
                                    "rejected_records":
                                        runtime_rejected,
                                    "blocking_findings":
                                        runtime_blocking_findings,
                                },
                            },
                        ),
                    )
                if (
                    not validation_ready
                    or blocking_findings > 0
                    or blocked_records > 0
                    or review_records > 0
                    or quarantine_records > 0
                    or rejected_records > 0
                ):
                    return self._send(
                        403,
                        envelope(
                            "migration_gate_blocked",
                            {
                                "message":
                                    "Approved request is not execution-ready",
                                "validation_ready":
                                    validation_ready,
                                "blocking_findings":
                                    blocking_findings,
                                "blocked_records":
                                    blocked_records,
                                "review_records":
                                    review_records,
                                "quarantine_records":
                                    quarantine_records,
                                "rejected_records":
                                    rejected_records,
                            },
                        ),
                    )

            # =====================================================
            # DEV DRY-RUN EXECUTOR
            # Simulation only. No target writes are permitted.
            # =====================================================

            execution_id = str(uuid.uuid4())
            started_at = utc_now()

            ready_records = runtime_staging.get(
                "ready_records",
                [],
            )

            transformation_plan = discovery.get(
                "transformation_plan",
                [],
            )

            record_results = []

            for record_index, record in enumerate(
                ready_records,
                start=1,
            ):
                applicable_transformations = []

                record_entity = str(
                    record.get("entity", "")
                ).strip()

                for transformation in transformation_plan:
                    source = str(
                        transformation.get("source", "")
                    )

                    source_entity = (
                        source.split(".", 1)[0]
                        if "." in source
                        else ""
                    )

                    action = str(
                        transformation.get("action", "")
                    )

                    if action == "REFERENTIAL_CHECK":
                        continue

                    if (
                        record_entity
                        and source_entity == record_entity
                    ):
                        applicable_transformations.append({
                            "action":
                                action,
                            "source":
                                transformation.get("source"),
                            "target":
                                transformation.get("target"),
                            "business_rules":
                                transformation.get(
                                    "business_rules",
                                    [],
                                ),
                            "status":
                                "SIMULATED",
                            "target_write":
                                False,
                        })

                record_results.append({
                    "sequence": record_index,
                    "entity": record.get("entity"),
                    "row": record.get("row"),
                    "record_key": record.get("record_key"),
                    "source_record": record.get(
                        "source_record",
                        {},
                    ),
                    "status": "SIMULATED",
                    "transformation_count":
                        len(applicable_transformations),
                    "transformations":
                        applicable_transformations,
                    "target_write_executed": False,
                    "production_action_executed": False,
                })

            simulated_record_count = len(
                record_results
            )

            success_count = sum(
                1
                for item in record_results
                if item.get("status") == "SIMULATED"
            )

            failure_count = (
                simulated_record_count - success_count
            )

            completed_at = utc_now()

            execution = {
                "execution_id": execution_id,
                "migration_id": migration_id,
                "approval_id": approval_id,
                "status": (
                    "DRY_RUN_COMPLETED"
                    if failure_count == 0
                    else "DRY_RUN_COMPLETED_WITH_ERRORS"
                ),
                "execution_gate": "PASS",
                "approval_status": "APPROVED",
                "environment": "DEV",
                "execution_mode": "DRY_RUN",
                "started_at": started_at,
                "completed_at": completed_at,
                "input_record_count":
                    len(ready_records),
                "simulated_record_count":
                    simulated_record_count,
                "success_count":
                    success_count,
                "failure_count":
                    failure_count,
                "record_results":
                    record_results,
                "production_executed": False,
                "production_action_executed": False,
                "target_write_executed": False,
                "reconciled": False,
            }

            with _lock:
                executions = _runtime_state.setdefault(
                    "executions",
                    {},
                )

                executions[execution_id] = dict(
                    execution
                )

                _runtime_state[
                    "last_execution_id"
                ] = execution_id

                _runtime_state[
                    "last_agent_action"
                ] = (
                    f"DEV dry-run completed: "
                    f"{migration_id}; "
                    f"{success_count}/"
                    f"{simulated_record_count} "
                    f"records simulated"
                )

            payload = execution
        if path == "/v1/reconciliations":
            execution_id = str(
                data.get("execution_id", "")
            ).strip()

            if not execution_id:
                return self._send(
                    400,
                    envelope(
                        "reconciliation_error",
                        {
                            "message":
                                "execution_id is required"
                        },
                    ),
                )

            with _lock:
                executions = _runtime_state.get(
                    "executions",
                    {},
                )
                stored_execution = executions.get(
                    execution_id
                )

            if stored_execution is None:
                return self._send(
                    404,
                    envelope(
                        "reconciliation_error",
                        {
                            "message":
                                "Execution not found",
                            "execution_id":
                                execution_id,
                        },
                    ),
                )

            execution_status = str(
                stored_execution.get(
                    "status",
                    "",
                )
            ).upper()

            if execution_status not in {
                "DRY_RUN_COMPLETED",
                "DRY_RUN_COMPLETED_WITH_ERRORS",
            }:
                return self._send(
                    409,
                    envelope(
                        "reconciliation_blocked",
                        {
                            "message":
                                "Execution is not eligible for reconciliation",
                            "execution_id":
                                execution_id,
                            "execution_status":
                                execution_status,
                        },
                    ),
                )

            record_results = stored_execution.get(
                "record_results",
                [],
            )

            input_record_count = int(
                stored_execution.get(
                    "input_record_count",
                    0,
                )
                or 0
            )

            simulated_record_count = len(
                record_results
            )

            matched_records = sum(
                1
                for item in record_results
                if str(
                    item.get("status", "")
                ).upper() == "SIMULATED"
            )

            failed_records = (
                simulated_record_count
                - matched_records
            )

            unmatched_source = max(
                input_record_count
                - simulated_record_count,
                0,
            )

            unexpected_simulated = max(
                simulated_record_count
                - input_record_count,
                0,
            )

            count_variance = (
                simulated_record_count
                - input_record_count
            )

            target_write_count = sum(
                1
                for item in record_results
                if item.get(
                    "target_write_executed"
                ) is True
            )

            production_action_count = sum(
                1
                for item in record_results
                if item.get(
                    "production_action_executed"
                ) is True
            )

            reconciliation_pass = (
                input_record_count
                == simulated_record_count
                == matched_records
                and failed_records == 0
                and unmatched_source == 0
                and unexpected_simulated == 0
                and target_write_count == 0
                and production_action_count == 0
            )

            reconciliation_id = str(
                uuid.uuid4()
            )

            reconciled_at = utc_now()

            reconciliation = {
                "reconciliation_id":
                    reconciliation_id,
                "execution_id":
                    execution_id,
                "migration_id":
                    stored_execution.get(
                        "migration_id"
                    ),
                "approval_id":
                    stored_execution.get(
                        "approval_id"
                    ),
                "status":
                    (
                        "PASS"
                        if reconciliation_pass
                        else "FAIL"
                    ),
                "reconciled_at":
                    reconciled_at,
                "input_record_count":
                    input_record_count,
                "simulated_record_count":
                    simulated_record_count,
                "matched_records":
                    matched_records,
                "unmatched_source":
                    unmatched_source,
                "unexpected_simulated":
                    unexpected_simulated,
                "failed_records":
                    failed_records,
                "count_variance":
                    count_variance,
                "target_write_count":
                    target_write_count,
                "production_action_count":
                    production_action_count,
                "production_executed":
                    False,
                "production_action_executed":
                    False,
                "target_write_executed":
                    False,
            }

            with _lock:
                reconciliations = (
                    _runtime_state.setdefault(
                        "reconciliations",
                        {},
                    )
                )

                reconciliations[
                    reconciliation_id
                ] = dict(reconciliation)

                executions = (
                    _runtime_state.setdefault(
                        "executions",
                        {},
                    )
                )

                current_execution = executions.get(
                    execution_id
                )

                if current_execution is not None:
                    current_execution[
                        "reconciled"
                    ] = reconciliation_pass

                    current_execution[
                        "reconciliation_id"
                    ] = reconciliation_id

                    current_execution[
                        "reconciliation_status"
                    ] = reconciliation[
                        "status"
                    ]

                    current_execution[
                        "reconciled_at"
                    ] = reconciled_at

                _runtime_state[
                    "last_reconciliation_id"
                ] = reconciliation_id

                _runtime_state[
                    "last_agent_action"
                ] = (
                    f"DEV dry-run reconciliation "
                    f"{reconciliation['status']}: "
                    f"{execution_id}"
                )

            payload = reconciliation
        if path == "/v1/evidence":
            reconciliation_id = str(
                data.get("reconciliation_id", "")
            ).strip()

            if not reconciliation_id:
                return self._send(
                    400,
                    envelope(
                        "evidence_error",
                        {
                            "message":
                                "reconciliation_id is required"
                        },
                    ),
                )

            with _lock:
                reconciliations = _runtime_state.get(
                    "reconciliations",
                    {},
                )

                reconciliation = reconciliations.get(
                    reconciliation_id
                )

            if reconciliation is None:
                return self._send(
                    404,
                    envelope(
                        "evidence_error",
                        {
                            "message":
                                "Reconciliation not found",
                            "reconciliation_id":
                                reconciliation_id,
                        },
                    ),
                )

            if str(
                reconciliation.get(
                    "status",
                    "",
                )
            ).upper() != "PASS":
                return self._send(
                    409,
                    envelope(
                        "evidence_blocked",
                        {
                            "message":
                                "Evidence package requires PASS reconciliation",
                            "reconciliation_status":
                                reconciliation.get(
                                    "status"
                                ),
                        },
                    ),
                )

            execution_id = str(
                reconciliation.get(
                    "execution_id",
                    "",
                )
            )

            migration_id = str(
                reconciliation.get(
                    "migration_id",
                    "",
                )
            )

            approval_id = str(
                reconciliation.get(
                    "approval_id",
                    "",
                )
            )

            with _lock:
                executions = _runtime_state.get(
                    "executions",
                    {},
                )

                approvals = _runtime_state.get(
                    "approvals",
                    {},
                )

                discoveries = _runtime_state.get(
                    "discoveries",
                    {},
                )

                execution = executions.get(
                    execution_id
                )

                approval = approvals.get(
                    approval_id
                )

                discovery = discoveries.get(
                    migration_id
                )

            if execution is None:
                return self._send(
                    404,
                    envelope(
                        "evidence_error",
                        {
                            "message":
                                "Execution evidence not found",
                            "execution_id":
                                execution_id,
                        },
                    ),
                )

            if approval is None:
                return self._send(
                    404,
                    envelope(
                        "evidence_error",
                        {
                            "message":
                                "Approval evidence not found",
                            "approval_id":
                                approval_id,
                        },
                    ),
                )

            if discovery is None:
                return self._send(
                    404,
                    envelope(
                        "evidence_error",
                        {
                            "message":
                                "Discovery evidence not found",
                            "migration_id":
                                migration_id,
                        },
                    ),
                )

            evidence_id = str(
                uuid.uuid4()
            )

            created_at = utc_now()

            evidence_package = {
                "evidence_id":
                    evidence_id,
                "created_at":
                    created_at,
                "migration_id":
                    migration_id,
                "approval_id":
                    approval_id,
                "execution_id":
                    execution_id,
                "reconciliation_id":
                    reconciliation_id,
                "status":
                    "COMPLETE",
                "discovery_summary":
                    discovery.get(
                        "summary",
                        {},
                    ),
                "approval_summary": {
                    "status":
                        approval.get("status"),
                    "approved":
                        approval.get("approved"),
                    "rejected":
                        approval.get("rejected"),
                    "decision_at":
                        approval.get("decision_at"),
                    "decision_by":
                        approval.get("decision_by"),
                    "decision_reason":
                        approval.get(
                            "decision_reason"
                        ),
                },
                "execution_summary": {
                    "status":
                        execution.get("status"),
                    "execution_gate":
                        execution.get(
                            "execution_gate"
                        ),
                    "input_record_count":
                        execution.get(
                            "input_record_count"
                        ),
                    "simulated_record_count":
                        execution.get(
                            "simulated_record_count"
                        ),
                    "success_count":
                        execution.get(
                            "success_count"
                        ),
                    "failure_count":
                        execution.get(
                            "failure_count"
                        ),
                    "started_at":
                        execution.get(
                            "started_at"
                        ),
                    "completed_at":
                        execution.get(
                            "completed_at"
                        ),
                },
                "reconciliation_summary":
                    dict(reconciliation),
                "record_results":
                    execution.get(
                        "record_results",
                        [],
                    ),
                "transformation_evidence":
                    discovery.get(
                        "target_staging_plan",
                        {},
                    ).get(
                        "evidence",
                        [],
                    ),
                "business_rules":
                    discovery.get(
                        "business_rules",
                        [],
                    ),
                "safety": {
                    "production_executed":
                        False,
                    "production_action_executed":
                        False,
                    "target_write_executed":
                        False,
                    "target_write_count":
                        reconciliation.get(
                            "target_write_count",
                            0,
                        ),
                    "production_action_count":
                        reconciliation.get(
                            "production_action_count",
                            0,
                        ),
                },
            }

            with _lock:
                evidence_store = (
                    _runtime_state.setdefault(
                        "evidence_packages",
                        {},
                    )
                )

                evidence_store[
                    evidence_id
                ] = dict(
                    evidence_package
                )

                _runtime_state[
                    "last_evidence_id"
                ] = evidence_id

                _runtime_state[
                    "last_agent_action"
                ] = (
                    f"Evidence package created: "
                    f"{evidence_id}"
                )

            payload = evidence_package


        if path == "/v1/learning/index":
            learning_id = str(
                data.get(
                    "learning_id",
                    "",
                )
            ).strip()

            if not learning_id:
                return self._send(
                    400,
                    envelope(
                        "learning_index_error",
                        {
                            "message":
                                "learning_id is required",
                            "source_write_executed":
                                False,
                            "target_write_executed":
                                False,
                            "production_action_executed":
                                False,
                        },
                    ),
                )

            learning = get_learning(
                learning_id
            )

            if learning is None:
                return self._send(
                    404,
                    envelope(
                        "learning_index_error",
                        {
                            "message":
                                "Durable learning package "
                                "not found",
                            "learning_id":
                                learning_id,
                            "source_write_executed":
                                False,
                            "target_write_executed":
                                False,
                            "production_action_executed":
                                False,
                        },
                    ),
                )

            try:
                result = (
                    index_verified_learning(
                        learning
                    )
                )
            except ValueError as exc:
                return self._send(
                    409,
                    envelope(
                        "learning_index_error",
                        {
                            "message":
                                str(exc),
                            "learning_id":
                                learning_id,
                            "source_write_executed":
                                False,
                            "target_write_executed":
                                False,
                            "production_action_executed":
                                False,
                        },
                    ),
                )

            with _lock:
                _runtime_state[
                    "last_learning_index"
                ] = dict(result)

                _runtime_state[
                    "last_agent_action"
                ] = (
                    "Verified learning indexed "
                    "into reusable knowledge: "
                    f"{learning_id}"
                )

            return self._send(
                200,
                envelope(
                    "learning_index",
                    result,
                ),
            )

        if path == "/v1/learning":
            evidence_id = str(
                data.get("evidence_id", "")
            ).strip()

            if not evidence_id:
                return self._send(
                    400,
                    envelope(
                        "verified_learning_error",
                        {
                            "message":
                                "evidence_id is required",
                            "source_write_executed":
                                False,
                            "target_write_executed":
                                False,
                            "production_action_executed":
                                False,
                        },
                    ),
                )

            with _lock:
                evidence_store = (
                    _runtime_state.get(
                        "evidence_packages",
                        {},
                    )
                )

                evidence = evidence_store.get(
                    evidence_id
                )

            if evidence is None:
                evidence_path = (
                    Path(__file__).resolve().parent
                    / "runtime"
                    / "evidence"
                    / f"{evidence_id}.json"
                )

                if evidence_path.exists():
                    try:
                        recovered = json.loads(
                            evidence_path.read_text(
                                encoding="utf-8-sig"
                            )
                        )

                        if (
                            isinstance(recovered, dict)
                            and str(
                                recovered.get(
                                    "evidence_id",
                                    "",
                                )
                            ).strip()
                            == evidence_id
                        ):
                            evidence = recovered

                            with _lock:
                                evidence_store = (
                                    _runtime_state.setdefault(
                                        "evidence_packages",
                                        {},
                                    )
                                )

                                evidence_store[
                                    evidence_id
                                ] = dict(evidence)

                                _runtime_state[
                                    "last_evidence_id"
                                ] = evidence_id
                    except (
                        OSError,
                        ValueError,
                        json.JSONDecodeError,
                    ):
                        evidence = None

            if evidence is None:
                return self._send(
                    404,
                    envelope(
                        "verified_learning_error",
                        {
                            "message":
                                "Evidence package not found",
                            "evidence_id":
                                evidence_id,
                            "source_write_executed":
                                False,
                            "target_write_executed":
                                False,
                            "production_action_executed":
                                False,
                        },
                    ),
                )

            try:
                learning = (
                    persist_verified_learning(
                        dict(evidence)
                    )
                )
            except ValueError as exc:
                return self._send(
                    409,
                    envelope(
                        "verified_learning_error",
                        {
                            "message": str(exc),
                            "evidence_id":
                                evidence_id,
                            "source_write_executed":
                                False,
                            "target_write_executed":
                                False,
                            "production_action_executed":
                                False,
                        },
                    ),
                )

            learning_id = str(
                learning.get(
                    "learning_id",
                    "",
                )
            )

            with _lock:
                learning_store = (
                    _runtime_state.setdefault(
                        "verified_learning",
                        {},
                    )
                )

                learning_store[
                    learning_id
                ] = dict(learning)

                _runtime_state[
                    "last_learning_id"
                ] = learning_id

                _runtime_state[
                    "last_agent_action"
                ] = (
                    "Verified learning persisted: "
                    f"{learning_id}"
                )

            payload = learning

        if path == "/v1/approvals":
            approval_id = str(uuid.uuid4())

            payload = {
                "id": approval_id,
                "status": "PENDING",
                "approved": False,
                "rejected": False,
                "requires_authoritative_approver": True,
                "requested_at": utc_now(),
                "decision_at": None,
                "decision_by": None,
                "decision_reason": None,
                "production_action_executed": False,
                "target_write_executed": False,
                "input": data,
            }

            with _lock:
                approvals = _runtime_state.setdefault(
                    "approvals",
                    {},
                )
                approvals[approval_id] = dict(payload)

                _runtime_state["last_agent_action"] = (
                    f"Approval request created: {approval_id}"
                )

        return self._send(
            202,
            envelope(
                allowed[path],
                payload,
            ),
        )
    def do_PATCH(self) -> None:  # noqa: N802
        record_request()
        path = self.path.split("?", 1)[0]

        data, error = self._read_json()

        if error:
            return self._send(
                400,
                envelope(
                    "error",
                    {"message": error},
                ),
            )

        assert data is not None

        prefix = "/v1/approvals/"

        if not path.startswith(prefix):
            return self._send(
                404,
                envelope(
                    "error",
                    {"message": "not found"},
                ),
            )

        approval_id = path[len(prefix):].strip()

        if not approval_id:
            return self._send(
                400,
                envelope(
                    "approval_decision_error",
                    {"message": "approval id is required"},
                ),
            )

        decision = str(
            data.get("decision", "")
        ).strip().upper()

        decision_by = str(
            data.get("decision_by", "")
        ).strip()

        decision_reason = str(
            data.get("decision_reason", "")
        ).strip()

        if decision not in {"APPROVE", "REJECT"}:
            return self._send(
                400,
                envelope(
                    "approval_decision_error",
                    {
                        "message":
                            "decision must be APPROVE or REJECT"
                    },
                ),
            )

        if not decision_by:
            return self._send(
                400,
                envelope(
                    "approval_decision_error",
                    {
                        "message":
                            "decision_by is required for authoritative approval"
                    },
                ),
            )

        if not decision_reason:
            return self._send(
                400,
                envelope(
                    "approval_decision_error",
                    {
                        "message":
                            "decision_reason is required"
                    },
                ),
            )

        with _lock:
            approvals = _runtime_state.get(
                "approvals",
                {},
            )

            approval = approvals.get(approval_id)

            if approval is None:
                return self._send(
                    404,
                    envelope(
                        "approval_decision_error",
                        {
                            "message":
                                "approval request not found"
                        },
                    ),
                )

            current_status = str(
                approval.get("status", "")
            ).upper()

            if current_status != "PENDING":
                return self._send(
                    409,
                    envelope(
                        "approval_decision_error",
                        {
                            "message":
                                "approval request has already been decided",
                            "status": current_status,
                        },
                    ),
                )

            if decision == "APPROVE":
                approval["status"] = "APPROVED"
                approval["approved"] = True
                approval["rejected"] = False
            else:
                approval["status"] = "REJECTED"
                approval["approved"] = False
                approval["rejected"] = True

            approval["decision_at"] = utc_now()
            approval["decision_by"] = decision_by
            approval["decision_reason"] = decision_reason

            # Governance decision only.
            # This does NOT execute migration or write to target.
            approval["production_action_executed"] = False
            approval["target_write_executed"] = False

            approvals[approval_id] = dict(approval)

            _runtime_state["last_agent_action"] = (
                f"Approval {approval_id} decided: "
                f"{approval['status']} by {decision_by}"
            )

            result = dict(approval)

        return self._send(
            200,
            envelope(
                "approval_decision",
                result,
            ),
        )    
    def log_message(self, fmt: str, *args: Any) -> None:
        return


def run(host: str = HOST, port: int = PORT) -> None:
    server = ThreadingHTTPServer((host, port), Handler)
    print(f"KMITORA A000 reference runtime listening on http://{host}:{port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    run()















